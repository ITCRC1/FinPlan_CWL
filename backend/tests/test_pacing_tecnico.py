# -*- coding: utf-8 -*-
"""El análisis técnico del año (`app/engine/pacing_tecnico.py`) y su Excel.

Con los XML reales del corte 30-sep-2026 el análisis dio: 5,786 RN / $5.40M
proyectados para 2027, 42% asegurado, esfuerzo 85% (3,349 RN que faltan contra
3,923 que 2026 recogió desde la misma fecha), reservas vs H&F 99.6% en los ocho
meses cerrados de 2026, y disponibilidad < sumar pickup < multiplicar por pace
(5,786 · 6,212 · 8,498). Acá se fijan las reglas con datos sintéticos.
"""
import io
from datetime import date, timedelta

from openpyxl import load_workbook

from app.engine import pacing as P
from app.engine.pacing_tecnico import analisis_tecnico
from app.export.pacing_tecnico_xlsx import construir_excel

CORTE = date(2026, 9, 30)


def _dias(desde, hasta, rn, rev, hist):
    out, d = {}, desde
    while d <= hasta:
        out[d.isoformat()] = [rn, rev, 30, 0, 0, 0, 1, rn * 2, 1 if hist else 0]
        d += timedelta(days=1)
    return out


def _snaps():
    t = {**_dias(date(2026, 1, 1), CORTE - timedelta(days=1), 20, 10000, True),
         **_dias(CORTE, date(2027, 12, 31), 6, 3000, False)}
    r = {k: [v[0], v[1] * 0.7, *v[2:]] for k, v in t.items()}
    base = {"as_of": CORTE.isoformat(), "date_from": "2026-01-01", "date_to": "2027-12-31", "has_forecast": True}
    return [{**base, "id": "t", "kind": "total", "days": t}, {**base, "id": "r", "kind": "rooms", "days": r}]


def _r(i, ins, arr, nts, rms=1, amt=500.0, st="A"):
    return {"id": str(i), "ins": ins, "arr": arr, "nts": nts, "rms": rms, "amt": amt * nts, "st": st,
            "fl": "", "ch": "C- AGENCIA", "blk": 0}


def _reservas():
    out = []
    # Enero 2026: 620 RN en H&F. 200 en libros al 30-sep-2025 (carga de migración del 1-sep), 418 después.
    out += [_r(f"a{k}", "2025-09-01", f"2026-01-{1 + k % 28:02d}", 2, rms=1) for k in range(100)]
    out += [_r(f"b{k}", "2025-12-15", f"2026-01-{1 + k % 28:02d}", 2, rms=1) for k in range(209)]
    # Enero 2027: 300 RN en libros hoy, todos con fecha posterior al año anterior.
    out += [_r(f"c{k}", "2026-08-01", f"2027-01-{1 + k % 28:02d}", 3, rms=1, amt=600) for k in range(100)]
    # Una reserva que toca el corte para que las reservas «cubran» la fecha del STLY.
    out += [_r("z", "2026-09-29", "2027-02-10", 1)]
    return out


def _A(**kw):
    return analisis_tecnico(_snaps(), _reservas(), {}, **kw)


def test_usa_la_misma_proyeccion_que_la_posicion():
    A = _A()
    assert A["anio"] == 2027 and A["escenario"] == "avail"
    ctx = P.Contexto(snaps=_snaps(), noches=P.expandir(_reservas()), config={})
    ref = P.con_en_sitio(ctx, P.construir(ctx, "total", 2027, "avail"), "total")
    assert abs(A["general"]["proj"] - ref["tot"]["proj"]) < 1e-6
    assert abs(A["general"]["proj_rev"] - ref["tot"]["proj_rev"]) < 1e-6
    for k in ("add", "mult"):
        alt = P.con_en_sitio(ctx, P.construir(ctx, "total", 2027, k), "total")
        assert abs(A["general"]["metodos"][k]["rn"] - alt["tot"]["proj"]) < 1e-6


def test_esfuerzo_y_asegurado_de_enero():
    m = _A()["meses"][0]
    # STLY = 200 (reservas al 30-sep-2025), cierre 2026 = 620, pickup del año anterior = 420
    assert m["stly"] == 200 and m["ly"] == 620 and m["ly_pick"] == 420
    assert m["otb"] == 6 * 31
    assert abs(m["falta"] - (m["proj"] - m["otb"])) < 1e-9
    assert abs(m["esfuerzo"] - m["falta"] / 420) < 1e-9
    assert abs(m["asegurado"] - m["otb"] / m["proj"]) < 1e-9
    assert m["comentario"].startswith("Enero 2027:")


def test_base_reconcilia_reservas_contra_history_forecast():
    b = _A()["base"]
    ene = b["filas"][0]
    assert ene["cerrado"] and ene["rn"] == 620
    assert abs(ene["reconcilia"] - 618 / 620) < 1e-9          # 309 reservas × 2 noches
    assert ene["stly"] == 200 and ene["pick"] == 420
    assert abs(ene["tasa"] - 420 / (930 - 200)) < 1e-9


def test_ventana_no_publica_lo_que_cae_antes_de_la_migracion():
    v = {x["i"]: x["puntos"] for x in _A()["base"]["ventana"]}
    # 1-ene-2026 menos 180 días cae antes de la carga del 1-sep-2025: no es observable
    assert v[0]["180"] is None
    # a 120 y 90 días (3-sep y 3-oct-2025) sólo estaba la migración
    assert abs(v[0]["120"] - 200 / 618) < 1e-9 and abs(v[0]["90"] - 200 / 618) < 1e-9
    assert abs(v[0]["30"] - 200 / 618) < 1e-9                    # al 2-dic: el resto entró el 15-dic


def test_textos_en_los_dos_idiomas():
    es, en = _A(lang="es"), _A(lang="en")
    assert es["meses"][0]["comentario"].startswith("Enero 2027")
    assert en["meses"][0]["comentario"].startswith("January 2027")
    assert es["general"]["porque"] and en["general"]["porque"]
    assert all(m["comentario"] for m in en["meses"])
    assert en["meses"][0]["confianza_txt"] in ("high", "medium", "low")


def test_escenario_sin_pickup_usa_disponibilidad():
    assert _A(escenario="otb")["escenario"] == "avail"


def test_el_excel_trae_las_seis_hojas_y_los_comentarios():
    A = P.limpiar(_A())
    wb = load_workbook(io.BytesIO(construir_excel(A, "Hotel Prueba")))
    assert wb.sheetnames == ["Resumen ejecutivo", "Análisis por mes", "Posición 2027", "Detalle por mes",
                             "Base 2026", "Métodos y glosario"]
    textos = [c.value for row in wb["Análisis por mes"].iter_rows() for c in row if isinstance(c.value, str)]
    assert A["meses"][0]["comentario"] in textos
    assert wb["Resumen ejecutivo"]["A1"].value == "Hotel Prueba · Análisis técnico del pacing 2027"
    en = load_workbook(io.BytesIO(construir_excel(P.limpiar(_A(lang="en")), "Hotel Prueba")))
    assert en.sheetnames[0] == "Executive summary"
