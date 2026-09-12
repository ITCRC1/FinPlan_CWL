"""Regla de ALLOCATION (queda ESCRITA y blindada): aplica a todo mes/año/escenario.

- 0220 (Employee Dining / comida): allocation TOTAL → su costo(5)/planilla(6)/opex(7)
  NO entran a los totales operativos (netean a 0 vía Distribución).
- 0161 (Laundry): SPLIT → la lavandería interna (planilla 6 + insumos 7) es allocation;
  el Laundry Services que vende afuera (ingreso 4 + costo de venta 5) SE QUEDA operativo.

Si alguien cambia ALLOCATION_EXCLUDE y rompe esto, este test falla.

La ÚNICA excepción es el espejo del Pre-Cierre (owner, 2026-09-10: «los
allocations en overhead […] esta regla es solo para este tab»): ahí el gasto SÍ
entra, en su línea de overhead, porque esa pantalla existe para revisar y no se
puede permitir que se escape plata mientras la 4999 no está posteada. El P&L de
verdad no cambia. Lo blindan `test_precierre_ve_el_allocation_como_overhead` y
`test_el_modo_overhead_no_se_puede_pedir_desde_afuera`.
"""
import io
import openpyxl
from pathlib import Path
from app.importers.gl_detail_importer import (
    parse_gl_detail, allocation_en_overhead, ALLOCATION_EXCLUDE, ALLOC_EXCL_COST,
    ALLOC_EXCL_PAYROLL, ALLOC_EXCL_OPEX,
)


def _gl_book(rows):
    """rows: (clase_lbl, dept_name, account_code, account_name, valor_enero)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    # Fila 15 = rótulo del bloque · Fila 14 = mes (el parser detecta por ambos).
    ws.cell(row=15, column=5, value="Actual 2099")
    ws.cell(row=14, column=5, value="January")
    r = 17
    for cls_lbl, dept, code, name, jan in rows:
        ws.cell(row=r, column=1, value=cls_lbl)
        ws.cell(row=r, column=2, value=dept)
        ws.cell(row=r, column=3, value=code)
        ws.cell(row=r, column=4, value=name)
        ws.cell(row=r, column=5, value=jan)   # mes 1
        r += 1
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_allocation_rule_locked():
    """La regla exacta queda fijada en código."""
    assert ALLOCATION_EXCLUDE == {"0220": {"5", "6", "7"}, "0161": {"6", "7"}}
    assert ALLOC_EXCL_COST == {"0220"}
    assert ALLOC_EXCL_PAYROLL == {"0220", "0161"}
    assert ALLOC_EXCL_OPEX == {"0220", "0161"}


def test_allocation_applied_on_parse():
    data = _gl_book([
        ("Rev",  "Lavanderia",                  "4700", "Laundry Services rev", 100),
        ("Cost", "Lavanderia",                  "5603", "Laundry Services cost", 50),
        ("Pay",  "Lavanderia",                  "6000", "Laundry payroll",      200),
        ("Opex", "Lavanderia",                  "7320", "Laundry supplies",      80),
        ("Cost", "Cafeteria Empleados",         "5700", "Dining cost",          300),
        ("Pay",  "Cafeteria Empleados",         "6000", "Dining payroll",       150),
        ("Opex", "Cafeteria Empleados",         "7065", "Dining opex",           20),
        ("Pay",  "Departamento de Habitaciones", "6000", "Rooms payroll",       500),
    ])
    blk = parse_gl_detail(data)[0]
    def depts(key):
        return {r["dept_code"] for r in blk[key]}

    # 0220 (Cafeteria) — allocation TOTAL: nada en costo/planilla/opex
    assert "0220" not in depts("costs")
    assert "0220" not in depts("payroll")
    assert "0220" not in depts("opex")

    # 0161 (Laundry) — split: planilla/opex FUERA; ingreso/costo de venta DENTRO
    assert "0161" not in depts("payroll")
    assert "0161" not in depts("opex")
    assert "0161" in depts("revenue")
    assert "0161" in depts("costs")

    # Control: un depto operativo normal entra en planilla
    assert "0110" in depts("payroll")


# ── La excepción del Pre-Cierre ──────────────────────────────────────────────

class _Esc:
    def __init__(self, es_precierre):
        self.es_precierre = es_precierre


def test_el_modo_overhead_solo_lo_abre_el_espejo_del_precierre():
    """Ningún otro escenario —ni la ausencia de escenario— lo enciende."""
    assert allocation_en_overhead(_Esc(True)) is True
    assert allocation_en_overhead(_Esc(False)) is False
    assert allocation_en_overhead(None) is False       # sin destino explícito
    assert allocation_en_overhead(object()) is False   # escenario sin el campo


def test_precierre_ve_el_allocation_como_overhead():
    """Con `en_overhead` el gasto de 0220/0161 ENTRA, y el resto no se mueve."""
    data = _gl_book([
        ("Rev",  "Lavanderia",                  "4700", "Laundry Services rev", 100),
        ("Cost", "Lavanderia",                  "5603", "Laundry Services cost", 50),
        ("Pay",  "Lavanderia",                  "6000", "Laundry payroll",      200),
        ("Opex", "Lavanderia",                  "7320", "Laundry supplies",      80),
        ("Cost", "Cafeteria Empleados",         "5700", "Dining cost",          300),
        ("Pay",  "Cafeteria Empleados",         "6000", "Dining payroll",       150),
        ("Opex", "Cafeteria Empleados",         "7065", "Dining opex",           20),
        ("Pay",  "Departamento de Habitaciones", "6000", "Rooms payroll",       500),
    ])
    blk = parse_gl_detail(data, en_overhead=True)[0]
    def depts(key):
        return {r["dept_code"] for r in blk[key]}

    # Lo que en el P&L de verdad desaparece, acá se ve:
    assert "0220" in depts("costs")
    assert "0220" in depts("payroll")
    assert "0220" in depts("opex")
    assert "0161" in depts("payroll")
    assert "0161" in depts("opex")
    # Y no se descartó ni una fila por allocation.
    assert blk["skipped"].get("allocation", 0) == 0
    # El operativo normal sigue igual: el modo no reacomoda nada más.
    assert "0110" in depts("payroll")
    assert "0161" in depts("revenue") and "0161" in depts("costs")


def test_el_modo_overhead_no_se_puede_pedir_desde_afuera():
    """El modo lo decide el DESTINO, no quien llama.

    Si algún día se expone como parámetro de la ruta, cualquier curl podría
    meter el gasto de allocation al Actual y el P&L seguiría cuadrando consigo
    mismo — el modo de falla más caro del sistema. Por eso se deriva del
    escenario y por eso este test mira el código fuente.
    """
    src = Path("app/api/scenarios_api.py").read_text(encoding="utf-8")
    assert "parse_gl_detail(data, en_overhead=allocation_en_overhead(forced))" in src
    assert "en_overhead: bool = Query" not in src
    assert "en_overhead=True" not in src   # nadie lo fija a mano


def test_el_gasto_por_clase_tambien_ve_el_allocation_en_precierre():
    """Los dos caminos del GOP tienen que ver lo MISMO.

    Owner, 2026-09-10: *«debe verlos como overhead los allocations»*. El cuadro
    calcula el GOP por NATURALEZA (ingreso menos clases 5/6/7) y el motor por
    DEPARTAMENTO. Desde que el espejo trae el gasto de allocation, el motor lo
    veia y `gasto-por-clase` no: US$40.700,88 de agosto 2026 de descuadre entre
    dos numeros que son el mismo, avisado en un banner que aparecia por diseno.
    """
    from app.api.gasto_por_clase_api import _excluidos, EXCLUIR_DE_GASTO

    class _Esc:
        def __init__(self, es_precierre):
            self.es_precierre = es_precierre

    assert _excluidos(_Esc(True)) == set()          # Pre-Cierre: los ve
    assert _excluidos(_Esc(False)) == EXCLUIR_DE_GASTO
    assert _excluidos(None) == EXCLUIR_DE_GASTO     # sin escenario, la regla de siempre
    # Y la regla de siempre sigue siendo la de siempre.
    assert EXCLUIR_DE_GASTO == {"0220", "0161", "0162"}


# ── El crédito del reparto, cuando YA está posteado ──────────────────────────
#
# Agosto 2026 fue el primer mes en que Integrity posteó la distribución: el
# archivo trajo `4999-0220` por −US$24.321,92 y `4999-0161` por −US$3.092,77, y
# el mismo monto repartido a los departamentos (6025, 7310, 7685).
#
# El parser saltaba la contrapartida SIEMPRE. Fuera del Pre-Cierre eso es
# correcto y es el par que siempre funcionó: `ALLOCATION_EXCLUDE` saca el gasto
# y el salto saca el crédito. Pero en el espejo del Pre-Cierre el gasto SÍ entra
# — y si el crédito no, la plata se cuenta dos veces: una en la línea de
# overhead del departamento de reparto y otra dentro de cada departamento que
# recibió su parte. En agosto eran US$27.414,69 de gasto inflado y la misma
# utilidad neta de menos.

BLOQUE = "Actual 2026"
DEPTOS = {"0110": "Habitaciones", "0161": "Lavanderia", "0220": "Cafeteria"}


def _plantilla(filas):
    """El libro REAL de FinPlan — el mismo que arma el Pre-Cierre para el espejo.

    `_gl_book` no sirve para esto: con filas armadas a mano la detección de
    columnas por contenido elige mal la del NOMBRE, y la contrapartida se
    reconoce justamente por el texto («distribu»).
    """
    from app.export.detail_excel import build_detail_workbook
    accts = [{"clase": clase, "grupo": "X", "dept_code": dept, "cuenta": cta,
              "nombre": nombre, "vals": {(BLOQUE, 1): float(monto)}, "orden": None}
             for dept, cta, nombre, clase, monto in filas]
    return build_detail_workbook([BLOQUE], accts, {}, DEPTOS)


FILAS_CON_REPARTO = [
    ("0220", "5420", "COST OF FOOD CAFETERIA - STAFF", "Cost", 12943.90),
    ("0220", "4999", "Expense Distribution", "Revenue", -24321.92),
    ("0161", "7320", "Laundry supplies", "Opex", 2133.55),
    ("0161", "4999", "Expense Distribution", "Revenue", -3092.77),
    ("0110", "7065", "Cleaning Supplies", "Opex", 1000.00),
]


def _por_cuenta(blk):
    return {(k, r["account_code"]): sum(r["months"].values())
            for k in ("revenue", "costs", "opex", "belowgop")
            for r in blk.get(k, [])}


def test_el_credito_de_reparto_no_entra_en_la_subida_normal():
    """Sin Pre-Cierre: ni el gasto del depto de reparto ni su crédito."""
    blk = parse_gl_detail(_plantilla(FILAS_CON_REPARTO))[0]
    fuera = _por_cuenta(blk)
    assert not [k for k in fuera if k[1] == "4999"], (
        f"la contrapartida no puede entrar en una subida normal: {sorted(fuera)}")
    assert ("costs", "5420") not in fuera, "y el gasto del 0220 tampoco"
    assert ("opex", "7065") in fuera, "un depto normal sí entra"


def test_el_credito_de_reparto_netea_en_pre_cierre():
    """Con `en_overhead`: entra, y entra como GASTO para netear la misma línea.

    En `revenue` cuadraría el total y estaría en la sección equivocada: su regla
    de mapeo apunta a OH_CAFETERIA / OH_LAUNDRY, que son líneas de gasto.
    """
    blk = parse_gl_detail(_plantilla(FILAS_CON_REPARTO), en_overhead=True)[0]
    fuera = _por_cuenta(blk)

    assert not [k for k in fuera if k[0] == "revenue" and k[1] == "4999"], (
        "el crédito de reparto NO puede quedar como ingreso negativo")
    creditos = {k for k in fuera if k[1] == "4999"}
    assert creditos == {("opex", "4999")}, f"tiene que ir a gasto: {creditos}"
    # `_por_cuenta` junta los dos departamentos en una llave; el total es el par.
    total = sum(sum(r["months"].values()) for r in blk["opex"]
                if r["account_code"] == "4999")
    assert round(total, 2) == -27414.69

    # Y el gasto bruto de los dos deptos de reparto entra, como siempre en este
    # modo: es lo que el crédito netea.
    assert ("costs", "5420") in fuera
    assert ("opex", "7320") in fuera


def test_el_departamento_del_credito_se_conserva():
    """Cada crédito tiene que quedar en SU departamento: si los dos cayeran en
    uno solo, netearían la línea equivocada y la otra quedaría inflada."""
    blk = parse_gl_detail(_plantilla(FILAS_CON_REPARTO), en_overhead=True)[0]
    por_dept = {r["dept_code"]: sum(r["months"].values())
                for r in blk["opex"] if r["account_code"] == "4999"}
    assert round(por_dept["0220"], 2) == -24321.92
    assert round(por_dept["0161"], 2) == -3092.77
