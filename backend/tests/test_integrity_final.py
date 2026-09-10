# -*- coding: utf-8 -*-
"""
EL TRADUCTOR TIENE QUE DAR LO MISMO QUE EL EXCEL DEL OWNER.

`app/importers/integrity_final.py` reemplaza las fórmulas que hoy se escriben a
mano cada mes en la hoja `Final` del libro de cierre. La prueba de que son
equivalentes no es leerlas: es correr el archivo de un mes ya cerrado y validado
—julio 2026— y comprobar que salen los mismos números.

El fixture `integrity_final_JUL2026.xlsx` trae SÓLO las columnas A–U, que es lo
que Integrity entrega. Que la prueba pase con eso demuestra que el traductor no
depende de ninguna de las columnas que el owner agrega a mano.
"""
import io
import json
import pathlib
from collections import defaultdict
from decimal import Decimal as D

import pytest

from app.importers import integrity_final as m

BASE = pathlib.Path(__file__).resolve().parents[1]
FIXTURE = BASE / "tests" / "fixtures" / "integrity_final_JUL2026.xlsx"
SEMILLA = BASE / "app" / "seed_data" / "CWL" / "mapd_integrity.json"

#: El tipo de cambio de julio 2026, tal como está en `Final!AD11` del libro.
TC_JULIO = D("454.75")

@pytest.fixture(scope="module")
def catalogo() -> dict:
    """El catálogo de departamentos de FinPlan, sembrado desde las constantes del
    motor. Es de donde sale la clasificación — el puente sólo traduce códigos."""
    from app.engine import pl_engine
    from app.seed_department_catalog import build_rows
    filas = build_rows()
    pl_engine.set_dept_catalog([{"dept_code": f["dept_code"],
                                 "default_pl_group": f.get("default_pl_group", ""),
                                 "parent_dept_code": f.get("parent_dept_code")}
                                for f in filas])
    return {f["dept_code"]: f for f in filas}


@pytest.fixture(scope="module")
def grupo_de(catalogo):
    """Grupo del P&L según FinPlan. `None` cuando el catálogo lo deja vacío a
    propósito (los que se abren POR CUENTA, como 280 Misceláneos): ahí no se
    inventa un grupo."""
    from app.engine import pl_engine

    def resolver(destino: str):
        fila = catalogo.get(destino)
        if fila is not None and not (fila.get("default_pl_group") or ""):
            return None
        return pl_engine.group_for_dept(destino)
    return resolver


#: Los grupos que FinPlan considera overhead. Sale del motor, no de una lista
#: escrita a mano acá — si el motor cambia, la prueba se entera.
@pytest.fixture(scope="module")
def overhead() -> set:
    from app.engine import pl_engine
    return set(pl_engine.OVERHEAD_DEPT_GROUPS)


@pytest.fixture(scope="module")
def mapd() -> dict:
    d = json.loads(SEMILLA.read_text(encoding="utf-8"))
    return {x["codigo"]: x for x in d["departamentos"]}


@pytest.fixture(scope="module")
def julio(mapd, grupo_de) -> dict:
    return m.leer(FIXTURE.read_bytes(), TC_JULIO, mapd, grupo_de)


@pytest.fixture(scope="module")
def por_categoria(julio) -> dict:
    acc = defaultdict(D)
    for f in julio["filas"]:
        acc[f["categoria"]] += f["mes_usd"]
    return acc


# ═════════════════════════════════════════════════════════════════════════════
# LA PRUEBA DE ORO — julio 2026, un mes ya cerrado y revisado por el owner
# ═════════════════════════════════════════════════════════════════════════════

def test_ingresos_totales(por_categoria):
    assert abs(por_categoria["Ingresos"] - D("248437.33")) < D("0.02")


def test_gasto_operativo(julio, overhead):
    """Los departamentos operativos, sin overhead y sin la clase 8."""
    t = sum((f["mes_usd"] for f in julio["filas"]
             if f["categoria"] not in ("Ingresos", "No Operativo")
             and f["grupo"] not in overhead), D(0))
    assert abs(t - D("147248.79")) < D("0.02")


def test_overhead(julio, overhead):
    t = sum((f["mes_usd"] for f in julio["filas"]
             if f["categoria"] not in ("Ingresos", "No Operativo")
             and f["grupo"] in overhead), D(0))
    assert abs(t - D("178789.87")) < D("0.02")


def test_utilidad_neta(por_categoria):
    """La línea de abajo. Si ésta cuadra, la cascada entera cuadra."""
    gastos = sum((por_categoria[c] for c in
                  ("Costo de Ventas", "Nómina", "Otros Gastos", "No Operativo",
                   "Allocation")), D(0))
    assert abs((por_categoria["Ingresos"] - gastos) - D("-94182.95")) < D("0.02")


def test_la_clase_8_completa_es_la_suma_de_sus_cinco_baldes(por_categoria):
    """La verificación parte la clase 8 en no operativo · capital · financieros ·
    depreciación · impuesto. Juntos tienen que dar el total de la clase."""
    baldes = D("18664.70") + D("9937.63") + D("0") + D("28343.41") + D("-40364.12")
    assert abs(por_categoria["No Operativo"] - baldes) < D("0.02")


# ═════════════════════════════════════════════════════════════════════════════
# LAS COLISIONES DE CÓDIGO — por qué el puente tiene que existir
#
# Los dos sistemas NO comparten espacio de códigos. Mapear por igualdad de
# código movería plata de departamento sin que ningún total cambie.
# ═════════════════════════════════════════════════════════════════════════════

def test_el_0130_de_integrity_es_un_restaurante_no_el_spa(julio):
    """Integrity `0130` = Restaurant Terra Kitchen. FinPlan `0130` = Spa
    (gerencia). Son $13.121,49 de A&B que irían al Spa si se mapeara por código."""
    filas = [f for f in julio["filas"] if f["depto"] == "0130"]
    assert filas, "el fixture de julio tiene que traer el 0130"
    assert all(f["destino_finplan"] == "0120" for f in filas)
    assert all(f["grupo"] == "FB" for f in filas)
    assert abs(sum((f["mes_usd"] for f in filas), D(0)) - D("13121.49")) < D("0.02")


def test_el_0151_de_integrity_es_el_gift_shop_no_la_tienda(julio):
    """FinPlan tiene DOS locales: `0151` Tienda y `0165` Gift Shop. El `0151` de
    Integrity es el segundo — sus cuentas dicen RETAIL GIFT SHOP."""
    filas = [f for f in julio["filas"] if f["depto"] == "0151"]
    assert filas
    assert all(f["destino_finplan"] == "0165" for f in filas)
    assert all(f["grupo"] == "RETAIL" for f in filas)


def test_el_private_bar_no_queda_tapado_dentro_de_ayb(julio):
    """`0128` son cuentas que dicen PRIVATE BAR y FinPlan lo tiene como centro de
    utilidad propio. El Excel del owner lo mandaba a A&B."""
    filas = [f for f in julio["filas"] if f["depto"] == "0128"]
    assert filas
    assert all(f["destino_finplan"] == "0121" for f in filas)
    assert all(f["grupo"] == "PRIVATE_BAR" for f in filas)


def test_innoceana_no_se_cuenta_dentro_de_tours(julio):
    """El Excel decía «Tours» en la columna de división, pero su propia hoja de
    revisión ya mostraba Innoceana como fila aparte."""
    filas = [f for f in julio["filas"] if f["depto"] == "0155"]
    assert filas
    assert all(f["grupo"] == "INNOCEANA" for f in filas)


def test_la_lavanderia_son_dos_departamentos(julio):
    """Diseño del owner: uno lleva el ingreso y el otro los gastos."""
    ingreso = [f for f in julio["filas"] if f["depto"] == "0160"]
    gasto = [f for f in julio["filas"] if f["depto"] == "0161"]
    assert all(f["destino_finplan"] == "0162" and f["grupo"] == "LAUNDRY" for f in ingreso)
    assert all(f["destino_finplan"] == "0161" and f["grupo"] == "LAUNDRY_OPS" for f in gasto)


def test_la_lavanderia_de_gastos_cierra_en_cero(julio):
    """`4999-0161 LAUNDRY EXPENSE DISTRIBUTION` reparte todo su costo a los
    departamentos que lo consumen: el departamento neto tiene que dar cero."""
    t = sum((f["mes_usd"] for f in julio["filas"] if f["depto"] == "0161"), D(0))
    assert abs(t) < D("0.02")


def test_un_destino_sin_grupo_no_se_inventa(julio):
    """`280` Misceláneos se abre POR CUENTA y el catálogo lo deja sin grupo a
    propósito. Poner el fallback rotularía como overhead lo que es ingreso."""
    filas = [f for f in julio["filas"] if f["destino_finplan"] == "280"]
    assert filas
    assert all(f["grupo"] is None for f in filas)


def test_todo_destino_existe_en_el_catalogo_de_finplan(mapd, catalogo):
    """Un destino que FinPlan no conoce dejaría la plata sin clasificar."""
    faltan = {d["codigo"]: d["destino_finplan"] for d in mapd.values()
              if d["destino_finplan"] not in catalogo}
    assert not faltan, f"destinos que no existen en FinPlan: {faltan}"


def test_ningun_departamento_quedo_sin_mapeo(julio):
    """Si aparece uno, su plata no llega a ningún departamento."""
    assert julio["sin_mapeo"] == [], (
        "departamentos sin mapear: "
        + str([(x["depto"], float(x["mes_usd"])) for x in julio["sin_mapeo"]]))


# ═════════════════════════════════════════════════════════════════════════════
# LAS REGLAS, UNA POR UNA
# ═════════════════════════════════════════════════════════════════════════════

@pytest.mark.parametrize("cuenta,esperado", [
    ("4000-0110", -1),      # ingreso: baja como crédito, se da vuelta
    ("4999-0110", 1),       # allocation: ya viene con el signo bueno
    ("5000-0120", 1),
    ("7000-0180", 1),
    ("8000-0240", 1),
])
def test_signo_por_clase(cuenta, esperado):
    assert m.signo_de(cuenta) == esperado


def test_el_signo_del_mes_lo_decide_el_acumulado():
    """Integrity da el mes en valor absoluto: sin mirar el acumulado, un ingreso
    y un gasto del mismo monto son indistinguibles."""
    assert m.monto_mes("4000-0110", 100, -500) == D("100")    # ingreso, se invierte
    assert m.monto_mes("4000-0110", 100, 500) == D("-100")
    assert m.monto_mes("7000-0180", 100, 500) == D("100")     # gasto, tal cual


@pytest.mark.parametrize("cuenta,esperado", [
    ("4000-0110", 4000),        # depto: SÍ se mapea
    ("4000", None),             # clase sola: no
    ("4000-0110-001", None),    # subdetalle: ya está sumado en el padre
    ("", None),
    ("no-es-cuenta", None),
])
def test_solo_las_cuentas_de_dos_segmentos_entran(cuenta, esperado):
    """La trampa que contaría la plata tres veces."""
    assert m.cuenta_base(cuenta) == esperado


def test_el_subdetalle_no_se_suma_dos_veces(julio):
    """Ninguna fila mapeada puede tener dos guiones."""
    assert not [f for f in julio["filas"] if f["cuenta"].count("-") != 1]


@pytest.mark.parametrize("cuenta,esperado", [
    ("4999-0110", "Allocation"), ("4000-0110", "Ingresos"),
    ("5000-0120", "Costo de Ventas"), ("6000-0110", "Nómina"),
    ("7000-0180", "Otros Gastos"), ("8000-0240", "No Operativo"),
    ("4000-0110-001", ""),
])
def test_categoria_usali(cuenta, esperado):
    assert m.categoria_usali(cuenta) == esperado


def test_depto_se_saca_de_la_cuenta():
    assert m.depto_de("4000-0110") == "0110"
    assert m.depto_de("4000-0110-001") == ""


# ── El tipo de cambio ────────────────────────────────────────────────────────

def test_el_tipo_de_cambio_es_obligatorio(mapd):
    """No se deduce del archivo. Inventarlo pondría todo el P&L a un TC que
    nadie decidió, y cuadraría igual."""
    with pytest.raises(ValueError):
        m.leer(FIXTURE.read_bytes(), None, mapd)


@pytest.mark.parametrize("tc", [0, -1, D("0")])
def test_un_tipo_de_cambio_invalido_no_pasa(tc):
    with pytest.raises(ValueError):
        m.a_dolares(D("1000"), tc)


def test_la_conversion_es_la_del_libro():
    """Fila 16 de julio: 47.911.719,75 colones / 454,75 = US$105.358,37."""
    crc = m.monto_mes("4000-0110", 47911719.75, -47911719.75)
    assert abs(m.a_dolares(crc, TC_JULIO) - D("105358.37")) < D("0.01")


# ── Números «sumables y legibles» ────────────────────────────────────────────

@pytest.mark.parametrize("crudo,esperado", [
    ("1,234.50", D("1234.50")),
    ("(1,234.50)", D("-1234.50")),      # negativo entre paréntesis
    ("$ 1,234.50", D("1234.50")),
    ("", D("0")), (None, D("0")), ("-", D("0")),
    (1234.5, D("1234.5")),
])
def test_texto_a_numero(crudo, esperado):
    assert m._dec(crudo) == esperado


# ── El formato falla ruidoso ─────────────────────────────────────────────────

def test_una_hoja_sin_los_encabezados_no_se_adivina():
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = "Final"
    wb.active["A1"] = "otra cosa"
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(m.FormatoInesperado) as e:
        m.leer(buf.getvalue(), TC_JULIO, {})
    assert "Mes Actual" in str(e.value)


def test_un_archivo_sin_la_hoja_Final_dice_que_hojas_trae():
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = "Otra"
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(m.FormatoInesperado) as e:
        m.leer(buf.getvalue(), TC_JULIO, {})
    assert "Otra" in str(e.value)


def test_los_encabezados_se_ubican_por_texto_no_por_posicion(julio):
    """Si Integrity agrega un renglón de título, buscar por número de fila leería
    la columna de al lado y el total seguiría cuadrando."""
    cols = julio["columnas"]
    assert cols["fila_encabezado"] == 10      # fila 11 en julio 2026
    assert cols["cuenta"] == 3                # D
    assert cols["mes"] == 17                  # R
    assert cols["acumulado"] == 20            # U


def test_un_depto_sin_mapeo_se_reporta_con_su_monto(mapd):
    """No se adivina y no se descarta en silencio: se dice cuánta plata es."""
    recortado = {k: v for k, v in mapd.items() if k != "0110"}
    r = m.leer(FIXTURE.read_bytes(), TC_JULIO, recortado)
    faltan = {x["depto"] for x in r["sin_mapeo"]}
    assert "0110" in faltan
    assert any(x["mes_usd"] != 0 for x in r["sin_mapeo"] if x["depto"] == "0110")


# ─── El archivo real no siempre viene como el libro del owner ─────────────────
#
# Todo lo de acá salió de subir el cierre de agosto 2026 por la pantalla: cada
# caso es un error que el usuario vio en pantalla, no uno imaginado.

def _libro(hojas: dict[str, list[list]]) -> bytes:
    """Un .xlsx armado a mano. `hojas` = {nombre: filas}."""
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nombre, filas in hojas.items():
        ws = wb.create_sheet(nombre)
        for f in filas:
            ws.append(f)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


CAB = ["Cuenta", "Descripción", "Mes Actual", "Acumulado"]
FILAS_OK = [["4000-0110", "Habitaciones", "1,000.00", "-5,000.00"]]
PUENTE_MIN = {"0110": {"nombre_integrity": "Hab", "destino_finplan": "0110"}}


def test_un_xls_viejo_dice_que_hacer_en_vez_de_reventar():
    """⚠️ Sin esto el 500 llega al navegador como «Failed to fetch».

    `openpyxl` sólo abre ZIP; un `.xls` es OLE2 y tira `BadZipFile`. Al subir
    sin manejar, la respuesta se corta y el browser no ve cabeceras CORS:
    reporta un fallo de red y el usuario no tiene forma de saber que el
    problema es su archivo. Pasó con el cierre de agosto 2026.
    """
    xls = bytes.fromhex("d0cf11e0a1b11ae1") + b"\x00" * 600
    with pytest.raises(m.FormatoInesperado) as e:
        m.leer(xls, TC_JULIO, PUENTE_MIN)
    assert ".xls" in str(e.value) and ".xlsx" in str(e.value)


def test_lo_que_no_es_un_libro_se_distingue_de_uno_cortado():
    """Un PDF y un .xlsx truncado no se arreglan igual: hay que decir cuál es."""
    with pytest.raises(m.FormatoInesperado) as pdf:
        m.leer(b"%PDF-1.7" + b"\x00" * 600, TC_JULIO, PUENTE_MIN)
    assert "no es un libro de Excel" in str(pdf.value)

    with pytest.raises(m.FormatoInesperado) as cortado:
        m.leer(b"PK\x03\x04" + b"\x00" * 40, TC_JULIO, PUENTE_MIN)
    assert "descarga cortada" in str(cortado.value)


def test_la_hoja_se_elige_por_contenido_y_no_por_nombre():
    """Integrity la entrega como `Sheet1`, no como `Final`.

    Exigir el nombre dejaba el mes afuera con «El archivo no trae la hoja
    Final» — un error que sólo se resuelve renombrando a mano una hoja que
    siempre se va a llamar igual.
    """
    r = m.leer(_libro({"Sheet1": [CAB] + FILAS_OK}), TC_JULIO, PUENTE_MIN)
    assert r["hoja"] == "Sheet1"
    assert len(r["filas"]) == 1


def test_la_cola_de_filas_vacias_no_molesta():
    """El archivo baja con centenares de filas vacías debajo del último dato."""
    hoja = [CAB] + FILAS_OK + [[None, None, None, None]] * 200
    r = m.leer(_libro({"Sheet1": hoja}), TC_JULIO, PUENTE_MIN)
    assert len(r["filas"]) == 1
    assert r["sin_cuenta"] == []      # vacío no es «fila sin cuenta»


def test_se_saltan_las_hojas_vacias_y_las_de_resumen():
    """La hoja buena es la que trae encabezados Y cuentas, esté donde esté."""
    libro = {"Sheet1": [[None] * 4] * 30,        # vacía
             "Resumen": [CAB],                    # encabezados, cero cuentas
             "Detalle": [CAB] + FILAS_OK}         # la buena
    assert m.leer(_libro(libro), TC_JULIO, PUENTE_MIN)["hoja"] == "Detalle"


def test_final_se_prefiere_cuando_existe():
    """Compatibilidad con el libro del owner: si está, manda."""
    libro = {"Sheet1": [CAB] + FILAS_OK,
             "Final": [CAB] + [["4000-0120", "F&B", "9.00", "-9.00"]]}
    r = m.leer(_libro(libro), TC_JULIO, {})
    assert r["hoja"] == "Final"
    assert r["filas"][0]["cuenta"] == "4000-0120"


def test_un_libro_entero_vacio_da_un_error_legible():
    with pytest.raises(m.FormatoInesperado) as e:
        m.leer(_libro({"Sheet1": [[None] * 4] * 20}), TC_JULIO, PUENTE_MIN)
    assert "Ninguna hoja" in str(e.value)


def test_el_rotulo_no_marca_la_columna_del_monto():
    """⚠️ El bug más caro que tuvo este módulo: los ingresos salían NEGATIVOS.

    En el archivo de agosto 2026 el rótulo «Acumulado» está una columna a la
    izquierda de sus montos —una celda combinada corrida—. El módulo leía la
    columna del rótulo, que está vacía; `_dec("")` es cero; y `monto_mes()`
    decide el signo MIRANDO EL ACUMULADO, así que con cero no entraba en la
    rama del crédito y a cada ingreso le quedaba `signo = -1`.

    Resultado: «VENTAS TOTALES −277.777». El P&L cuadraba consigo mismo, no
    hubo ningún error, y el mes entero estaba dado vuelta.
    """
    cab = [None] * 21
    cab[2], cab[9], cab[17], cab[19] = "Cuenta", "Descripción", "Mes Actual", "Acumulado"

    def fila(cuenta, desc, mes, acum):
        f = [None] * 21
        f[3], f[9], f[17], f[20] = cuenta, desc, mes, acum   # monto en 20, no 19
        return f

    libro = _libro({"Sheet1": [cab,
                               fila("4000-0110", "ROOMS", 1000.0, -5000.0),
                               fila("7065-0110", "LIMPIEZA", 200.0, 800.0)]})
    r = m.leer(libro, TC_JULIO, PUENTE_MIN)

    assert r["columnas"]["acumulado"] == 20, "se quedó en la columna del rótulo"
    ingreso = next(f for f in r["filas"] if f["cuenta"] == "4000-0110")
    assert ingreso["mes_usd"] > 0, "un ingreso no puede salir negativo"
    gasto = next(f for f in r["filas"] if f["cuenta"] == "7065-0110")
    assert gasto["mes_usd"] > 0


def test_sin_columna_de_acumulado_se_niega_a_adivinar():
    """Sin acumulado no se puede saber el signo. Antes leía cero y seguía."""
    cab = [None] * 21
    cab[2], cab[17], cab[19] = "Cuenta", "Mes Actual", "Acumulado"
    f = [None] * 21
    f[3], f[17] = "4000-0110", 1000.0          # sin acumulado en ninguna parte
    with pytest.raises(m.FormatoInesperado) as e:
        m.leer(_libro({"Sheet1": [cab, f]}), TC_JULIO, PUENTE_MIN)
    assert "signo" in str(e.value)
