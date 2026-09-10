# -*- coding: utf-8 -*-
"""Lee el estado de resultados CRUDO de Integrity — la hoja `Final`.

## Qué es este archivo

Lo que Integrity entrega cada mes: una hoja `Final` con el mayor del período.
**Del sistema baja sólo hasta la columna U.** Todo lo que está a la derecha en el
libro del owner —el signo, el tipo de cambio, el departamento, la división y la
categoría USALI— lo agrega él a mano, mes a mes, con fórmulas.

Este módulo es esas fórmulas. Cada función lleva en su docstring **la fórmula de
Excel que reemplaza**, copiada del libro (`Conc JUL 2026.xlsx`, julio 2026), para
que la equivalencia se pueda verificar leyendo, no confiando.

## Por qué es puro

No toca la base, no escribe nada, no sabe qué es un escenario. Recibe bytes y
devuelve filas. Eso es lo que permite probarlo con casos armados a mano y correr
la prueba de oro —julio 2026, que ya está cerrado y validado— sin levantar nada.

## Las dos trampas

⚠️ **El subdetalle NO se mapea.** Una cuenta de Integrity viene en tres niveles:
`4000` (clase), `4000-0110` (clase + departamento) y `4000-0110-001` (el
movimiento). El libro mapea **sólo las de dos segmentos**: son el total del
departamento, y las de tres ya están sumadas adentro. Contar las tres cuenta la
plata tres veces. Es la razón de ser de `cuenta_base()`.

⚠️ **El signo de los ingresos viene invertido.** En el mayor un ingreso es un
crédito y baja negativo, salvo `4999`, que son las contrapartidas de allocation y
ya vienen con el signo que corresponde.
"""
from __future__ import annotations

from decimal import Decimal

#: El nombre con el que venía la hoja en el libro del owner. Es una
#: PREFERENCIA, no un requisito: la hoja se elige por contenido en
#: `_elegir_hoja()` — Integrity también la entrega como `Sheet1`.
HOJA = "Final"

#: Los encabezados que tiene que traer, y con los que se ubican las columnas.
#: Se buscan POR TEXTO y no por letra: si Integrity mueve una columna, esto falla
#: diciendo qué no encontró en vez de leer la de al lado en silencio.
ENCABEZADOS = {"cuenta": "Cuenta", "descripcion": "Descripción",
               "mes": "Mes Actual", "acumulado": "Acumulado"}

#: Primer dígito de la cuenta → categoría USALI.
CATEGORIA_POR_CLASE = {"4": "Ingresos", "5": "Costo de Ventas", "6": "Nómina",
                       "7": "Otros Gastos", "8": "No Operativo"}

#: La clase que son contrapartidas de reparto, no ingreso.
CLASE_ALLOCATION = "4999"

ZERO = Decimal("0")


def signo_de(cuenta: str) -> int:
    """Excel: ``IF(LEFT(D,4)="4999",1,IF(LEFT(D,1)="4",-1,1))``

    Un ingreso baja del mayor como crédito, o sea negativo, y hay que darlo
    vuelta para leerlo. `4999` es la excepción: son las contrapartidas de
    allocation, que ya vienen con el signo bueno — invertirlas haría que un
    reparto sumara en vez de restar y el total seguiría cuadrando.
    """
    c = (cuenta or "").strip()
    if c[:4] == CLASE_ALLOCATION:
        return 1
    return -1 if c[:1] == "4" else 1


def monto_mes(cuenta: str, mes_crudo, acumulado_crudo) -> Decimal:
    """Excel: ``signo * IF($U<0,-$R,$R)``  (columna X, «Mes ₡ con signo»)

    ⚠️ El signo del MES lo decide el ACUMULADO, no el mes. Integrity entrega la
    columna del mes en valor absoluto y la del acumulado con signo; sin mirar el
    acumulado, un ingreso y un gasto del mismo monto son indistinguibles.
    """
    r, u = _dec(mes_crudo), _dec(acumulado_crudo)
    return signo_de(cuenta) * (-r if u < ZERO else r)


def monto_acumulado(cuenta: str, acumulado_crudo) -> Decimal:
    """Excel: ``signo * $U``  (columna Y, «Acumulado ₡ con signo»)"""
    return signo_de(cuenta) * _dec(acumulado_crudo)


def a_dolares(monto_crc: Decimal, tc: Decimal) -> Decimal:
    """Excel: ``$X/$AD$11``  (columnas Z y AA)

    El TC **no tiene default**. Cambia todos los meses y es un dato del cierre,
    no una constante del sistema: inventarlo pondría todo el P&L a un tipo de
    cambio que nadie decidió, y cuadraría igual.
    """
    if tc is None or Decimal(str(tc)) <= ZERO:
        raise ValueError("El tipo de cambio es obligatorio y tiene que ser mayor que cero.")
    return monto_crc / Decimal(str(tc))


def cuenta_base(cuenta: str) -> int | None:
    """Excel: ``IF(AND($D<>"",LEN($D)-LEN(SUBSTITUTE($D,"-",""))=1),VALUE(LEFT($D,4)),"")``

    Devuelve la clase de cuenta **sólo para las filas de dos segmentos**
    (`4000-0110`). Para la clase sola (`4000`) y para el movimiento
    (`4000-0110-001`) devuelve `None`, que es lo que las deja fuera de los
    totales. Ver la trampa del subdetalle arriba.
    """
    c = (cuenta or "").strip()
    if not c or c.count("-") != 1:
        return None
    try:
        return int(c[:4])
    except ValueError:
        return None


def depto_de(cuenta: str) -> str:
    """Excel: ``MID($D,6,4)``  (columna AF)"""
    return "" if cuenta_base(cuenta) is None else (cuenta or "").strip()[5:9]


def categoria_usali(cuenta: str) -> str:
    """Excel: ``IF($AE=4999,"Allocation",IF(LEFT($D,1)="4","Ingresos", … ))``  (AH)"""
    if cuenta_base(cuenta) is None:
        return ""
    c = (cuenta or "").strip()
    if c[:4] == CLASE_ALLOCATION:
        return "Allocation"
    return CATEGORIA_POR_CLASE.get(c[:1], "")


def _dec(v) -> Decimal:
    """Un número sumable. Integrity entrega la columna como texto con separadores
    de miles y, a veces, el negativo entre paréntesis — el «convertir a números
    sumables y legibles» que hoy se pide a mano."""
    if v is None or v == "":
        return ZERO
    if isinstance(v, Decimal):
        return v
    if isinstance(v, (int, float)):
        return Decimal(str(v))
    s = str(v).strip().replace(" ", "").replace(" ", "")
    negativo = s.startswith("(") and s.endswith(")")
    if negativo:
        s = s[1:-1]
    s = s.replace("$", "").replace("₡", "").replace(",", "")
    if not s or s in ("-", "."):
        return ZERO
    try:
        d = Decimal(s)
    except Exception:
        return ZERO
    return -d if negativo else d


# ─── Lectura de la hoja ───────────────────────────────────────────────────────

class FormatoInesperado(ValueError):
    """La hoja no tiene la forma que este módulo sabe leer. El mensaje dice QUÉ
    faltó: es la diferencia entre arreglarlo en un minuto y abrir el Excel a
    adivinar."""


def _texto(v) -> str:
    return "" if v is None else str(v).strip()


def _parece_cuenta(v) -> bool:
    """`4000`, `4000-0110`, `4000-0110-001`. El formato de Integrity."""
    s = _texto(v)
    if not s:
        return False
    partes = s.split("-")
    return partes[0].isdigit() and len(partes[0]) == 4 and all(p.isdigit() for p in partes[1:])


def _tiene_encabezados(filas: list[tuple]) -> bool:
    """¿Esta hoja trae la fila con «Mes Actual» y «Acumulado»?"""
    for fila in filas:
        textos = {_texto(v).lower() for v in fila}
        if (ENCABEZADOS["mes"].lower() in textos
                and ENCABEZADOS["acumulado"].lower() in textos):
            return True
    return False


def _elegir_hoja(wb) -> tuple[str, list[tuple]]:
    """Cuál de las hojas es el estado de resultados, y sus filas.

    ⚠️ **No se elige por nombre.** El módulo nació exigiendo que se llamara
    «Final», que es como venía el libro del owner. Pero Integrity la entrega
    como `Sheet1` —y cualquier «Guardar como» de Excel también—, así que el
    nombre no identifica nada: es el rótulo que quedó, no un dato del formato.

    Se elige igual que la fila de encabezados y la columna de cuenta: **por
    contenido**. La hoja del estado de resultados es la que trae «Mes Actual» y
    «Acumulado» juntos. `Final` se prueba primero por si el libro trae varias
    hojas con forma parecida y una es la buena de siempre.
    """
    orden = ([HOJA] if HOJA in wb.sheetnames else []) +             [n for n in wb.sheetnames if n != HOJA]
    #: La que tiene encabezados PERO ninguna cuenta. Se guarda como segunda
    #: opción: si al final no hay ninguna completa, devolverla deja que
    #: `ubicar_encabezados` explique con precisión qué le falta, en vez de un
    #: «ninguna hoja sirve» que no ayuda a nadie.
    a_medias = None
    for nombre in orden:
        filas = list(wb[nombre].iter_rows(values_only=True))
        if not _tiene_encabezados(filas):
            continue
        if any(_parece_cuenta(v) for f in filas for v in f):
            return nombre, filas
        if a_medias is None:
            a_medias = (nombre, filas)
    if a_medias is not None:
        return a_medias
    raise FormatoInesperado(
        f"Ninguna hoja del libro trae «{ENCABEZADOS['mes']}» y "
        f"«{ENCABEZADOS['acumulado']}» en la misma fila. "
        f"Hojas: {', '.join(wb.sheetnames)}. "
        f"¿Es el estado de resultados de Integrity?")


def _columna_de_montos(datos: list[tuple], desde: int, hasta: int) -> int | None:
    """La columna con más montos entre `desde` y `hasta`, o `None` si no hay.

    ⚠️ **El rótulo no marca la columna del número.** En el archivo de agosto
    2026 «Acumulado» está en la columna 19 y los montos en la 20 —una celda
    combinada corrida—, y el módulo leía la 19: vacía. Como `_dec("")` es cero
    y `monto_mes()` decide el SIGNO mirando el acumulado, cada ingreso salía
    NEGATIVO. El P&L cuadraba consigo mismo y el total de ventas daba menos
    277.777, sin un solo error.

    Es la misma trampa que la columna de cuenta, y se resuelve igual: contando
    contenido. Los montos también vienen como texto con separadores de miles,
    así que cuenta las dos formas.
    """
    conteo: dict[int, int] = {}
    for f in datos:
        for j in range(desde, min(hasta + 1, len(f))):
            v = f[j]
            if isinstance(v, (int, float, Decimal)):
                if v != 0:
                    conteo[j] = conteo.get(j, 0) + 1
            elif isinstance(v, str) and _dec(v) != ZERO:
                conteo[j] = conteo.get(j, 0) + 1
    return max(conteo, key=lambda j: conteo[j]) if conteo else None


def ubicar_encabezados(filas: list[tuple]) -> dict:
    """Dónde están las columnas, buscadas POR TEXTO.

    La fila de encabezados es la que trae «Mes Actual» y «Acumulado» juntos —en
    el archivo de julio 2026 es la 11, pero no se asume: si Integrity agrega un
    renglón de título, buscar por posición leería la columna equivocada y el
    total seguiría cuadrando.

    ⚠️ La columna de la CUENTA no se toma del encabezado. En el archivo real el
    rótulo «Cuenta» está en `C` y los códigos en `D` (celdas combinadas). Se
    detecta por CONTENIDO: la columna con más valores que parecen cuenta.
    """
    for i, fila in enumerate(filas):
        textos = {_texto(v).lower(): j for j, v in enumerate(fila)}
        col_mes = textos.get(ENCABEZADOS["mes"].lower())
        col_acum = textos.get(ENCABEZADOS["acumulado"].lower())
        if col_mes is None or col_acum is None:
            continue
        datos = filas[i + 1:]
        conteo: dict[int, int] = {}
        for f in datos:
            for j, v in enumerate(f):
                if _parece_cuenta(v):
                    conteo[j] = conteo.get(j, 0) + 1
        if not conteo:
            raise FormatoInesperado(
                f"Se encontró la fila de encabezados ({i + 1}) pero ninguna columna "
                f"trae códigos de cuenta con el formato de Integrity (`4000-0110`).")
        col_cuenta = max(conteo, key=lambda j: conteo[j])
        # Los rótulos dan el punto de partida; el número manda. La ventana de
        # «Mes Actual» termina donde empieza el rótulo del acumulado, para que
        # una no se coma la columna de la otra.
        tope_mes = col_acum - 1 if col_acum > col_mes else col_mes + 3
        mes = _columna_de_montos(datos, col_mes, tope_mes)
        acumulado = _columna_de_montos(datos, col_acum, col_acum + 3)
        if mes is None or acumulado is None or mes == acumulado:
            raise FormatoInesperado(
                f"Se ubicó la fila de encabezados ({i + 1}) pero no las dos "
                f"columnas de monto: «{ENCABEZADOS['mes']}» resolvió a "
                f"{mes} y «{ENCABEZADOS['acumulado']}» a {acumulado}. "
                f"Sin el acumulado no se puede saber el signo de un ingreso.")
        return {"fila_encabezado": i, "cuenta": col_cuenta, "mes": mes,
                "acumulado": acumulado,
                "descripcion": textos.get(ENCABEZADOS["descripcion"].lower())}
    raise FormatoInesperado(
        f"La hoja no trae una fila con «{ENCABEZADOS['mes']}» y "
        f"«{ENCABEZADOS['acumulado']}». ¿Es el estado de resultados de Integrity?")


def mapear_filas(filas: list[tuple], cols: dict, tc, puente: dict,
                 grupo_de=None) -> dict:
    """Las filas ya traducidas, más los departamentos que nadie mapeó.

    `puente` = {codigo_integrity: {nombre_integrity, destino_finplan}} — la
    semilla `seed_data/<HOTEL>/mapd_integrity.json`.

    `grupo_de(destino)` devuelve el grupo del P&L al que pertenece ese
    departamento **según FinPlan**, o `None` si el catálogo lo deja vacío a
    propósito (los departamentos que se abren POR CUENTA, como `280`
    Misceláneos). Cuando es `None` no se inventa un grupo: la línea la resuelve
    el mapeo de cuentas, y poner el fallback —`OTHER_OVERHEAD`— rotularía como
    overhead lo que en realidad es ingreso.

    ⚠️ **La clasificación no sale de este archivo, sale del catálogo de
    FinPlan.** El puente sólo traduce el CÓDIGO, porque los dos sistemas no
    comparten espacio de códigos: Integrity `0130` es el Restaurant Terra
    Kitchen y FinPlan `0130` es «Spa (gerencia)». Ver la nota de la semilla.

    Un departamento que no está en el puente **no se adivina**: va a `sin_mapeo`
    con su monto, para que se pueda decidir si es ruido o si falta media
    operación. Es la misma regla que el importador de Channel Mix aplica con los
    market codes sin canal.
    """
    tc = Decimal(str(tc)) if tc is not None else None
    salida, sin_mapeo = [], {}
    #: Filas con monto y SIN código de cuenta reconocible. No se descartan en
    #: silencio: así se perdieron $40.613 del gasto de Habitaciones en el Actual
    #: 2024, en dos renglones, y el descuadre apareció meses después.
    sin_cuenta: list[dict] = []
    #: El subdetalle (`4000-0110-001`), agrupado por su cuenta padre de dos
    #: segmentos. No entra a los totales —el padre ya lo suma— pero se conserva
    #: para poder comprobar que efectivamente suma.
    subdetalle: dict[str, Decimal] = {}
    #: El nivel de POSICIÓN de las cuentas 6 (`6000-0111-501`), fila por fila.
    #: No entra a ningún total —el padre ya lo suma—: es el mismo dinero,
    #: abierto por quién lo cobra.
    posiciones: list[dict] = []
    for n, f in enumerate(filas[cols["fila_encabezado"] + 1:], cols["fila_encabezado"] + 2):
        cuenta = _texto(f[cols["cuenta"]] if cols["cuenta"] < len(f) else "")
        crudo_m = f[cols["mes"]] if cols["mes"] < len(f) else None
        if not _parece_cuenta(cuenta):
            desc = _texto(f[cols["descripcion"]]
                          if cols["descripcion"] is not None
                          and cols["descripcion"] < len(f) else "")
            # ⚠️ Sólo cuenta como «fila sin cuenta» si trae DESCRIPCIÓN. El
            # reporte de Integrity lleva sus propios subtotales —«UTILIDAD
            # OPERATIVA», «UTILIDAD/PÉRDIDA NETA»— que tienen monto y no tienen
            # código, y son correctos: su rótulo va en la columna de etiquetas,
            # no en la de detalle. Sin este filtro, todos los meses aparecerían
            # dos hallazgos criticos falsos y la lista se aprenderia a ignorar.
            if _dec(crudo_m) != ZERO and desc:
                sin_cuenta.append({"fila": n, "monto_crc": _dec(crudo_m),
                                   "texto": desc})
            continue
        base = cuenta_base(cuenta)
        if base is None:
            # Clase sola (`4000`) o subdetalle (`4000-0110-001`). No entra a los
            # totales: el padre ya lo suma. El subdetalle se guarda aparte.
            if cuenta.count("-") == 2:
                padre = cuenta.rsplit("-", 1)[0]
                m_crc = monto_mes(cuenta, crudo_m, f[cols["acumulado"]]
                                  if cols["acumulado"] < len(f) else None)
                subdetalle[padre] = subdetalle.get(padre, ZERO) + m_crc
                # ⚠️ En las cuentas 6 este nivel ES LA POSICIÓN.
                #
                # `6000-0111-501` = concepto 6000, depto 0111, posición 501
                # (Front Desk Agent) — CLAUDE.md §12.1. Owner, 2026-09-10: *«la
                # posición en las cuentas 6 es el tercer nivel»*.
                #
                # Hasta acá el subdetalle se sumaba por cuenta padre y el
                # código de posición se tiraba, así que la planilla real solo
                # se podía mirar por departamento. Se conserva la fila entera
                # para poder abrirla por posición.
                #
                # Verificado sobre agosto 2026: las 116 cuentas de planilla
                # suman EXACTAMENTE lo mismo por posición que a nivel de
                # cuenta. No es un total nuevo: es el mismo, abierto.
                if cuenta[:1] == "6" and m_crc != ZERO:
                    # El departamento sale del PADRE: `depto_de` lee una cuenta
                    # de dos segmentos, y ésta tiene tres.
                    dep = depto_de(padre)
                    mm = puente.get(dep)
                    posiciones.append({
                        "fila": n,
                        "cuenta": cuenta,
                        "cuenta_base": cuenta_base(cuenta.rsplit("-", 1)[0]),
                        "posicion": cuenta.rsplit("-", 1)[1],
                        "depto": dep,
                        "destino_finplan": (mm or {}).get("destino_finplan", ""),
                        "descripcion": _texto(
                            f[cols["descripcion"]]
                            if cols["descripcion"] is not None
                            and cols["descripcion"] < len(f) else ""),
                        "mes_crc": m_crc,
                        "mes_usd": a_dolares(m_crc, tc),
                    })
            continue
        crudo_acum = f[cols["acumulado"]] if cols["acumulado"] < len(f) else None
        mes_crc = monto_mes(cuenta, crudo_m, crudo_acum)
        acum_crc = monto_acumulado(cuenta, crudo_acum)
        depto = depto_de(cuenta)
        m = puente.get(depto)
        destino = (m or {}).get("destino_finplan", "")
        fila = {
            "fila": n,
            "cuenta": cuenta,
            "cuenta_base": base,
            "depto": depto,
            "descripcion": _texto(f[cols["descripcion"]]
                                  if cols["descripcion"] is not None
                                  and cols["descripcion"] < len(f) else ""),
            "categoria": categoria_usali(cuenta),
            "destino_finplan": destino,
            "grupo": (grupo_de(destino) if grupo_de and destino else None),
            "nombre_integrity": (m or {}).get("nombre_integrity", ""),
            "mes_crc": mes_crc,
            "acumulado_crc": acum_crc,
            "mes_usd": a_dolares(mes_crc, tc),
            "acumulado_usd": a_dolares(acum_crc, tc),
        }
        salida.append(fila)
        if m is None:
            acc = sin_mapeo.setdefault(depto, {"depto": depto, "mes_usd": ZERO,
                                               "cuentas": []})
            acc["mes_usd"] += fila["mes_usd"]
            acc["cuentas"].append(cuenta)
    return {"filas": salida,
            "sin_mapeo": sorted(sin_mapeo.values(), key=lambda x: -abs(x["mes_usd"])),
            "sin_cuenta": sin_cuenta,
            "subdetalle": {k: v / tc for k, v in subdetalle.items()},
            "posiciones": posiciones}


#: Firma de los `.xls` viejos (OLE2). No son ZIP, y `openpyxl` sólo lee ZIP.
FIRMA_XLS_VIEJO = bytes.fromhex("d0cf11e0a1b11ae1")
#: Firma de cualquier ZIP — y por lo tanto de un `.xlsx` sano.
FIRMA_ZIP = b"PK"


def _por_que_no_abrio(data: bytes) -> str:
    """Por qué `openpyxl` no pudo abrirlo, dicho para quien subió el archivo.

    Se mira la firma en vez del nombre: el nombre lo elige la persona y miente
    —renombrar un `.xls` a `.xlsx` no lo convierte—, los primeros bytes no.
    """
    if data[:8] == FIRMA_XLS_VIEJO:
        return ("El archivo está en el formato viejo de Excel (.xls) y este "
                "lector sólo abre .xlsx. Abrilo en Excel y guardalo con "
                "«Guardar como → Libro de Excel (*.xlsx)». Renombrar la "
                "extensión no alcanza: es otro formato por dentro.")
    if data[:2] != FIRMA_ZIP:
        return ("El archivo no es un libro de Excel: no empieza como un .xlsx "
                "ni como un .xls. Puede ser un CSV, un PDF o una descarga "
                "incompleta.")
    return ("El archivo dice ser .xlsx pero no se pudo abrir. Suele ser una "
            "descarga cortada: bajalo de nuevo de Integrity y reintentá.")


def leer(data: bytes, tc, puente: dict, grupo_de=None) -> dict:
    """El archivo crudo de Integrity → filas traducidas. Es la puerta del módulo."""
    import io as _io
    import openpyxl
    if tc is None:
        raise ValueError("El tipo de cambio es obligatorio: no se deduce del archivo.")
    try:
        wb = openpyxl.load_workbook(_io.BytesIO(data), read_only=True,
                                    data_only=True)
    except FormatoInesperado:
        raise
    except Exception as e:
        # ⚠️ **Sin este try el 500 llega al navegador como «Failed to fetch».**
        # Cuando la excepción sube sin manejar, la respuesta se corta y el
        # browser no ve cabeceras CORS: reporta un fallo de red, no el 500. El
        # usuario ve un error que no dice nada y no tiene forma de saber que el
        # problema es su archivo. Pasó con un `.xls` de agosto 2026.
        raise FormatoInesperado(_por_que_no_abrio(data)) from e
    nombre_hoja, filas = _elegir_hoja(wb)
    cols = ubicar_encabezados(filas)
    return {**mapear_filas(filas, cols, tc, puente, grupo_de),
            "columnas": cols, "hoja": nombre_hoja}
