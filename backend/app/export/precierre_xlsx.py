# -*- coding: utf-8 -*-
"""Las descargas de Pre-Cierre. Todas editables, todas re-subibles.

Pedido del owner (2026-08-31): *«me gustaría la opción de bajar a Excel en el
mismo formato de FinPlan que usa para bajar, y todos los datos de este módulo
deben tener la opción de baja a Excel con formato editable»*.

Son tres, y cada una contesta otra pregunta:

| Descarga | Para qué |
|---|---|
| `plantilla_finplan` | **El formato estándar.** El mismo archivo que la app baja y vuelve a leer, ya lleno con el mes. Se edita y se sube por la puerta de siempre. |
| `filas_editables` | El detalle traducido, fila por fila, con de dónde salió cada número. Es lo que se mira cuando un total no convence. |
| `listado` | Las vueltas que lleva el mes, con quién subió qué y cuándo. |

**Editable de verdad: valores, no fórmulas.** Un Excel con fórmulas se ve igual
pero se rompe al editar una celda de la que otras dependen, y el que lo edita no
tiene forma de saberlo. Acá cada celda es un número.

⚠️ `plantilla_finplan` **no arma el layout**: se lo pide a
`export/detail_excel.build_detail_workbook`, que es el que genera la plantilla
oficial. Si el orden de las líneas cambia en la app, esta descarga cambia con
él. Copiarlo sería tener dos plantillas que se parecen hasta que dejan de
parecerse.
"""
from __future__ import annotations

import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.export.detail_excel import CLASE_BY_PREFIX, build_detail_workbook
from app.importers.verificacion import CONTROLES

ZERO = Decimal("0")

#: Control de la verificación → la clave de la hoja de revisión que lo alimenta.
#: Los once, verificados contra el bloque VERIF del libro del owner (julio 2026).
CONTROL_DESDE = {
    "VER_INGRESOS": "TOTAL INCOMES",
    "VER_GASTO_OPERATIVO": "Total Operationg expenses",
    "VER_OVERHEAD": "TOTAL OVERHEAD EXPENSES",
    "VER_GOP": "TOTAL GROSS OPERATING PROFIT",
    "VER_NO_OPERATIVO": "TOTAL Owners Expenses",
    "VER_EBITDA": "EBITDA",
    "VER_CAPITAL": "CAPITAL EXPENSE",
    "VER_FINANCIEROS": "FINANCIAL EXPENSES",
    "VER_DEPRECIACION": "TOTAL DEPRECIATIONS",
    "VER_IMPUESTO": "bg.Income Tax (30%)",
    "VER_UTILIDAD_NETA": "EARNINGS AFTER INCOME TAXES",
}

_HDR = PatternFill("solid", fgColor="16402A")
_BLANCO = Font(bold=True, color="FFFFFF", size=10)


def controles(valores: dict) -> dict[str, Decimal]:
    """Los once totales de control, desde la hoja de revisión ya calculada."""
    return {c.codigo: Decimal(str(valores.get(CONTROL_DESDE[c.codigo], 0)))
            for c in CONTROLES if c.codigo in CONTROL_DESDE}


def plantilla_finplan(mes: int, anio: int, etiqueta: str, filas: list[dict],
                      valores: dict, nombres_depto: dict[str, str],
                      linea_de=None) -> bytes:
    """El mes en el formato ESTÁNDAR de FinPlan — el que se baja y se sube.

    `linea_de(depto, cuenta)` devuelve la línea del P&L de esa cuenta, si se
    sabe. Es la columna «Línea del P&L» de la plantilla; vacía no rompe nada,
    pero es la que deja ver a dónde va cada cuenta sin abrir el mapeo.
    """
    accts = []
    for f in filas:
        cuenta = str(f["cuenta_base"] or "")
        if not cuenta:
            continue
        depto = f["destino_finplan"]
        accts.append({
            "clase": CLASE_BY_PREFIX.get(cuenta[:1], ""),
            "grupo": (linea_de(depto, cuenta) if linea_de else "") or "",
            "dept_code": depto,
            "cuenta": cuenta,
            "nombre": f["descripcion"],
            # Un solo bloque y un solo mes: es un cierre mensual, no un año.
            "vals": {(etiqueta, mes): float(f["mes_usd"])},
        })
    verif = {etiqueta: {codigo: {mes: float(monto)}
                        for codigo, monto in controles(valores).items()}}
    return build_detail_workbook([etiqueta], accts, {}, nombres_depto, verif)


def _hoja(wb: Workbook, titulo: str, encabezados: list[str], filas: list[list]):
    ws = wb.create_sheet(titulo) if wb.sheetnames != ["Sheet"] else wb.active
    ws.title = titulo
    for j, t in enumerate(encabezados, 1):
        c = ws.cell(1, j, t)
        c.fill = _HDR
        c.font = _BLANCO
        c.alignment = Alignment(horizontal="center")
    for i, fila in enumerate(filas, 2):
        for j, v in enumerate(fila, 1):
            celda = ws.cell(i, j, float(v) if isinstance(v, Decimal) else v)
            if isinstance(v, Decimal):
                celda.number_format = "#,##0.00"
    for j, t in enumerate(encabezados, 1):
        ancho = max([len(str(t))] + [len(str(f[j - 1])) for f in filas[:200]] or [10])
        ws.column_dimensions[ws.cell(1, j).column_letter].width = min(max(ancho + 2, 10), 46)
    ws.freeze_panes = "A2"
    return ws


def filas_editables(filas: list[dict]) -> bytes:
    """El detalle traducido, con de dónde salió cada número.

    Lleva la fila del Excel de origen: sin eso, revisar un número obliga a
    buscarlo a ojo en el archivo de Integrity.
    """
    wb = Workbook()
    _hoja(wb, "Detalle traducido",
          ["Fila origen", "Cuenta", "Cuenta base", "Nombre de cuenta",
           "Depto Integrity", "Depto FinPlan", "Grupo P&L", "Categoría USALI",
           "Mes US$", "Acumulado US$"],
          [[f["fila"], f["cuenta"], f["cuenta_base"], f["descripcion"],
            f["depto"], f["destino_finplan"], f["grupo"] or "(por cuenta)",
            f["categoria"], Decimal(str(f["mes_usd"])),
            Decimal(str(f["acumulado_usd"]))] for f in filas])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def listado(precierres: list[dict]) -> bytes:
    """Las vueltas del mes: quién subió qué, cuándo, y en qué quedó."""
    wb = Workbook()
    _hoja(wb, "Pre-cierres",
          ["Año", "Mes", "Estado", "Tipo de cambio", "Archivo", "Subido por",
           "Subido en", "Hallazgos abiertos"],
          [[p["anio"], p["mes"], p["estado"], Decimal(str(p["tc"])), p["archivo"],
            p["subido_por"], p["creado_en"] or "", p["hallazgos_abiertos"]]
           for p in precierres])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
