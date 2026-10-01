# -*- coding: utf-8 -*-
"""El análisis técnico del pacing en Excel, para la Junta.

Owner, 2026-09-30: *«que se pueda bajar a Excel en forma profesional con todos
los comentarios bien legibles, listo para compartir con la Junta»*.

Todo sale de `engine.pacing_tecnico.analisis_tecnico`: este módulo sólo
acomoda. Los textos vienen ya redactados del motor —los mismos que se ven en
pantalla— para que el Excel y la pantalla nunca digan cosas distintas.

⚠️ **Los comentarios van en bloques por mes, no en una columna.** Una columna de
comentarios al final de una tabla de 20 columnas obliga a imprimir en letra
diminuta o a cortar el texto. Cada mes es una ficha: cifras clave arriba,
comentario completo debajo, ancho de página.
"""
from __future__ import annotations

import math
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.engine.pacing_tecnico import MESES
from app.export.excel_base import C, workbook_to_bytes

NAVY, NAVY_MID, CLARO, GRIS = C["navy"], C["navy_mid"], C["blue_light"], C["gray_light"]
VERDE, AMBAR, ROJO = "1A7F5A", "B7791F", "C0392B"
F_RN, F_USD, F_PCT = "#,##0", '"$"#,##0', "0.0%"
_lin = Side(style="thin", color=C["border"])
BORDE = Border(left=_lin, right=_lin, top=_lin, bottom=_lin)

L = {
    "es": {
        "hojas": ["Resumen ejecutivo", "Análisis por mes", "Posición {T}", "Detalle por mes", "Base {B}", "Métodos y glosario"],
        "titulo": "{hotel} · Análisis técnico del pacing {T}",
        "sub": "Corte de datos: {corte} · Métrica: {kind} · Proyección: método {met} · Generado: {hoy}",
        "kinds": {"total": "Total Revenue", "rooms": "Room Revenue"},
        "met": {"avail": "disponibilidad", "add": "sumar pickup", "mult": "multiplicar por pace"},
        "kpi": "Indicadores clave", "ind": "Indicador", "valor": "Valor",
        "kpis": [("otb", "Noches en libros hoy (RN)", F_RN), ("occ_otb", "Ocupación en libros", F_PCT),
                 ("otb_rev", "Ingreso en libros", F_USD), ("stly", "En libros hace un año a la misma fecha (STLY)", F_RN),
                 ("pace", "Pace (reservas individuales vs STLY)", F_PCT), ("proj", "Proyectado (RN)", F_RN),
                 ("proj_occ", "Ocupación proyectada", F_PCT), ("proj_rev", "Ingreso proyectado", F_USD),
                 ("ly", "Cierre del año anterior (RN)", F_RN), ("ly_rev", "Ingreso del año anterior", F_USD),
                 ("asegurado", "Parte del proyectado ya asegurada", F_PCT), ("falta", "Pickup necesario (RN)", F_RN),
                 ("ly_pick", "Pickup del año anterior desde la misma fecha (RN)", F_RN),
                 ("esfuerzo", "Esfuerzo pedido (necesario / logrado el año anterior)", F_PCT)],
        "concl": "Conclusiones", "porque": "Por qué confiar en el proyectado", "limites": "Límites del análisis",
        "conf": "Confianza", "razones": "Por qué esta confianza: ",
        "ficha": [("otb", "En libros", F_RN), ("pace", "Pace", F_PCT), ("proj", "Proyectado", F_RN),
                  ("proj_occ", "Ocup. proy.", F_PCT), ("proj_rev", "Ingreso proy.", F_USD),
                  ("asegurado", "Asegurado", F_PCT), ("falta", "Faltan RN", F_RN), ("esfuerzo", "Esfuerzo", F_PCT)],
        "pos": ["Mes", "Disponibles", "En libros RN", "Ocup.", "Tarifa en libros", "Ingreso en libros", "STLY RN", "Pace",
                "Bloqueos", "Año ant. RN", "Base año ant.", "Pickup necesario", "Proyectado RN", "Ocup. proy.",
                "Ingreso proy.", "Estado"],
        "det": ["Mes", "En libros", "STLY", "Pace", "Proyectado", "Ocup. proy.", "Ingreso proy.", "vs año ant. RN",
                "vs año ant. ingreso", "Asegurado", "Faltan RN", "Pickup año ant.", "Esfuerzo", "RN por semana",
                "Ritmo últ. 30 días /sem", "Tarifa en libros", "Tarifa año ant. misma fecha", "Var. tarifa",
                "Noches libres", "Valor +5 pp ocup.", "Valor +5% tarifa", "Confianza"],
        "base": ["Mes", "Estado", "Disponibles", "RN", "Ocup.", "Ingreso", "Tarifa", "Presupuesto RN", "Presupuesto ingreso",
                 "vs presupuesto RN", "En libros al {st}", "Recogido después", "Tasa de llenado de libres", "Reservas vs H&F"],
        "real": "Real", "curso": "En curso", "cerrado_h": "Sin operación", "total": "Total",
        "ventana": "Ventana de reserva {B}: % de las noches del cierre que ya estaban reservadas X días antes del mes",
        "dias": "{d} días", "nd": "n/d",
        "nota_base": ("La «tasa de llenado de libres» es la parte de las habitaciones que estaban libres a la fecha de corte del año "
                      "anterior y que se vendieron después. Es el único supuesto de la proyección: el año {T} llenará la misma "
                      "parte de sus habitaciones libres. «n/d» en la ventana de reserva: la fecha cae antes de la carga inicial de "
                      "reservas y el dato no es observable."),
        "metodos": ["Mes", "Disponibilidad", "Sumar pickup", "Multiplicar por pace", "Usado", "Ingreso usado"],
        "glosario": "Cómo se calcula", "base_lab": {"real": "Real", "alcance": "Proyectado", "reservas": "Reservas", "sin": "—"},
        "estados": {"supera": "Supera año ant.", "rn_arriba_ingreso_abajo": "RN arriba / ingreso abajo", "cerca": "Cerca",
                    "debajo": "Debajo", "nuevo": "Nuevo", "sin_actividad": "Sin actividad"},
        "gl": [
            ("En libros (OTB)", "Noches reservadas hoy según el History & Forecast de Opera; incluye bloqueos de grupo."),
            ("STLY", "Same Time Last Year: lo que había en libros para el mismo mes hace exactamente un año."),
            ("Pace", "Reservas individuales de hoy divididas entre el STLY. Sobre 100% = se va más adelantado que el año anterior."),
            ("Bloqueos", "Noches de grupo en el History & Forecast que todavía no tienen reservas individuales (rooming list)."),
            ("Método disponibilidad", "Proyectado = en libros + tasa de llenado del año anterior × habitaciones libres. Nunca pasa de la capacidad."),
            ("Método sumar pickup", "Proyectado = en libros + las noches que el año anterior recogió desde la misma fecha."),
            ("Método multiplicar por pace", "Proyectado = en libros × (cierre del año anterior / STLY). Amplifica el pace: es el más agresivo."),
            ("Esfuerzo", "Pickup necesario para llegar al proyectado dividido entre el pickup que el año anterior logró desde la misma fecha."),
            ("Asegurado", "Parte del proyectado que ya está en libros."),
            ("Confianza", "Puntaje por mes: +2 si 60% o más está asegurado (+1 si 30%–60%); +2 si el esfuerzo es 90% o menos "
                          "(+1 si hasta 110%); +1 si la referencia del año anterior es un cierre real; +1 si el pace es 100% o más; "
                          "−1 o −2 si los bloqueos de grupo pasan del 25% o 50% de lo que está en libros; −1 si la tarifa de "
                          "referencia es atípica. Alta 5 o más, media 3–4, baja 2 o menos."),
            ("Ingreso", "Incluye lo reservado más una estimación del consumo en sitio (configurable en Pacing → Cargas)."),
        ],
    },
    "en": {
        "hojas": ["Executive summary", "Monthly analysis", "Position {T}", "Monthly detail", "{B} base", "Methods and glossary"],
        "titulo": "{hotel} · Technical pacing analysis {T}",
        "sub": "Data as of: {corte} · Metric: {kind} · Projection: {met} method · Generated: {hoy}",
        "kinds": {"total": "Total Revenue", "rooms": "Room Revenue"},
        "met": {"avail": "availability", "add": "add pickup", "mult": "multiply by pace"},
        "kpi": "Key indicators", "ind": "Indicator", "valor": "Value",
        "kpis": [("otb", "Room nights on the books today (RN)", F_RN), ("occ_otb", "Occupancy on the books", F_PCT),
                 ("otb_rev", "Revenue on the books", F_USD), ("stly", "On the books a year ago at the same date (STLY)", F_RN),
                 ("pace", "Pace (individual reservations vs STLY)", F_PCT), ("proj", "Projection (RN)", F_RN),
                 ("proj_occ", "Projected occupancy", F_PCT), ("proj_rev", "Projected revenue", F_USD),
                 ("ly", "Prior-year close (RN)", F_RN), ("ly_rev", "Prior-year revenue", F_USD),
                 ("asegurado", "Share of the projection already secured", F_PCT), ("falta", "Pickup needed (RN)", F_RN),
                 ("ly_pick", "Prior-year pickup from the same date (RN)", F_RN),
                 ("esfuerzo", "Effort required (needed / achieved last year)", F_PCT)],
        "concl": "Conclusions", "porque": "Why the projection can be trusted", "limites": "Limits of the analysis",
        "conf": "Confidence", "razones": "Why this confidence: ",
        "ficha": [("otb", "On the books", F_RN), ("pace", "Pace", F_PCT), ("proj", "Projection", F_RN),
                  ("proj_occ", "Proj. occ.", F_PCT), ("proj_rev", "Proj. revenue", F_USD),
                  ("asegurado", "Secured", F_PCT), ("falta", "RN to go", F_RN), ("esfuerzo", "Effort", F_PCT)],
        "pos": ["Month", "Available", "OTB RN", "Occ.", "OTB rate", "OTB revenue", "STLY RN", "Pace", "Blocks",
                "Prior year RN", "Prior-year basis", "Pickup needed", "Projected RN", "Proj. occ.", "Proj. revenue", "Status"],
        "det": ["Month", "On the books", "STLY", "Pace", "Projection", "Proj. occ.", "Proj. revenue", "vs prior year RN",
                "vs prior year revenue", "Secured", "RN to go", "Prior-year pickup", "Effort", "RN per week",
                "Last 30 days pace /wk", "OTB rate", "Prior-year rate same date", "Rate var.", "Unsold nights",
                "Value +5 occ. pts", "Value +5% rate", "Confidence"],
        "base": ["Month", "Status", "Available", "RN", "Occ.", "Revenue", "Rate", "Budget RN", "Budget revenue",
                 "vs budget RN", "On the books at {st}", "Picked up after", "Fill rate of free rooms", "Reservations vs H&F"],
        "real": "Actual", "curso": "Open", "cerrado_h": "No operation", "total": "Total",
        "ventana": "{B} booking window: % of the closing room nights already booked X days before the month",
        "dias": "{d} days", "nd": "n/a",
        "nota_base": ("The 'fill rate of free rooms' is the share of rooms that were free at last year's cut-off date and were sold "
                      "afterwards. It is the projection's only assumption: {T} will fill the same share of its free rooms. "
                      "'n/a' in the booking window: the date falls before the initial reservation load and is not observable."),
        "metodos": ["Month", "Availability", "Add pickup", "Multiply by pace", "Used", "Revenue used"],
        "glosario": "How it is calculated", "base_lab": {"real": "Actual", "alcance": "Projected", "reservas": "Reservations", "sin": "—"},
        "estados": {"supera": "Above prior year", "rn_arriba_ingreso_abajo": "RN up / revenue down", "cerca": "Close",
                    "debajo": "Below", "nuevo": "New", "sin_actividad": "No activity"},
        "gl": [
            ("On the books (OTB)", "Room nights booked today per Opera's History & Forecast; includes group blocks."),
            ("STLY", "Same Time Last Year: what was on the books for the same month exactly one year earlier."),
            ("Pace", "Today's individual reservations divided by STLY. Above 100% = ahead of last year."),
            ("Blocks", "Group room nights in the History & Forecast that do not have individual reservations yet (rooming list)."),
            ("Availability method", "Projection = on the books + last year's fill rate × free rooms. Never exceeds capacity."),
            ("Add pickup method", "Projection = on the books + the room nights last year picked up from the same date."),
            ("Multiply by pace method", "Projection = on the books × (last year's close / STLY). It amplifies the pace: the most aggressive."),
            ("Effort", "Pickup needed to reach the projection divided by the pickup last year achieved from the same date."),
            ("Secured", "Share of the projection already on the books."),
            ("Confidence", "Monthly score: +2 if 60% or more is secured (+1 if 30%–60%); +2 if effort is 90% or less "
                           "(+1 if up to 110%); +1 if the prior-year reference is an actual close; +1 if pace is 100% or more; "
                           "−1 or −2 if group blocks exceed 25% or 50% of the books; −1 if the reference rate is atypical. "
                           "High 5 or more, medium 3–4, low 2 or less."),
            ("Revenue", "Includes what is booked plus an estimate of on-site spend (configurable in Pacing → Uploads)."),
        ],
    },
}


def _f(sz=10, bold=False, color="1A1A2E", italic=False) -> Font:
    return Font(name="Calibri", size=sz, bold=bold, color=color, italic=italic)


def _fill(c) -> PatternFill:
    return PatternFill("solid", fgColor=c)


def _alto(texto: str, ancho_chars: float, base: float = 15.0) -> float:
    """Alto de fila para que un texto envuelto se lea completo (openpyxl no autoajusta)."""
    lineas = sum(max(1, math.ceil(len(p) / max(ancho_chars, 10))) for p in (texto or "").split("\n"))
    return max(base, 14.5 * lineas + 4)


def _encabezado(ws, fila: int, textos: list[str], anchos: list[float] | None = None) -> None:
    for j, t in enumerate(textos, start=1):
        c = ws.cell(row=fila, column=j, value=t)
        c.font = _f(10, True, "FFFFFF"); c.fill = _fill(NAVY_MID); c.border = BORDE
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[fila].height = 32
    if anchos:
        for j, w in enumerate(anchos, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w


def _celda(ws, fila, col, valor, fmt=None, bold=False, fill=None, color="1A1A2E", h="right", wrap=False):
    c = ws.cell(row=fila, column=col, value=valor)
    c.font = _f(10, bold, color); c.border = BORDE
    c.alignment = Alignment(horizontal=h, vertical="center", wrap_text=wrap)
    if fmt and isinstance(valor, (int, float)):
        c.number_format = fmt
    if fill:
        c.fill = _fill(fill)
    return c


def _titulo(ws, A, txt, ultima_col: int) -> int:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ultima_col)
    c = ws.cell(row=1, column=1, value=txt["titulo"])
    c.font = _f(15, True, "FFFFFF"); c.fill = _fill(NAVY); c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 30
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ultima_col)
    c = ws.cell(row=2, column=1, value=txt["sub"])
    c.font = _f(9, False, C["text_mid"], True); c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[2].height = 18
    return 4


def _seccion(ws, fila, texto, ultima_col) -> int:
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ultima_col)
    c = ws.cell(row=fila, column=1, value=texto)
    c.font = _f(12, True, NAVY); c.border = Border(bottom=Side(style="medium", color=NAVY))
    ws.row_dimensions[fila].height = 22
    return fila + 1


def _parrafo(ws, fila, texto, ultima_col, ancho, desde=1, bold=False, viñeta=True) -> int:
    ws.merge_cells(start_row=fila, start_column=desde, end_row=fila, end_column=ultima_col)
    c = ws.cell(row=fila, column=desde, value=("•  " if viñeta else "") + texto)
    c.font = _f(10.5, bold); c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
    ws.row_dimensions[fila].height = _alto(texto, ancho)
    return fila + 1


def _pagina(ws, horizontal=True, filas_titulo: str | None = None):
    ws.page_setup.orientation = "landscape" if horizontal else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.sheet_view.showGridLines = False
    ws.oddFooter.center.text = "&P / &N"
    if filas_titulo:
        ws.print_title_rows = filas_titulo


def _color_conf(n: str) -> str:
    return {"alta": VERDE, "media": AMBAR, "baja": ROJO}.get(n, C["text_mid"])


# ───────────────────────────── hojas ─────────────────────────────
def _resumen(ws, A, txt, lang):
    g = A["general"]
    ANCHO = [52, 20, 14, 14, 14, 14]
    for j, w in enumerate(ANCHO, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    total_chars = sum(ANCHO) * 1.05
    f = _titulo(ws, A, txt, 6)
    f = _seccion(ws, f, txt["kpi"], 6)
    _encabezado(ws, f, [txt["ind"], txt["valor"]]); f += 1
    datos = {**g, "ly": g["ly"], "ly_rev": g["ly_rev"]}
    for i, (k, et, fmt) in enumerate(txt["kpis"]):
        bg = GRIS if i % 2 else None
        _celda(ws, f, 1, et, h="left", fill=bg)
        _celda(ws, f, 2, datos.get(k), fmt, bold=k in ("proj", "proj_rev"), fill=bg)
        f += 1
    f += 1
    f = _seccion(ws, f, txt["concl"], 6)
    for c in g["conclusiones"]:
        f = _parrafo(ws, f, c, 6, total_chars)
    f += 1
    f = _seccion(ws, f, txt["porque"], 6)
    for k, r in enumerate(g["porque"], start=1):
        f = _parrafo(ws, f, f"{k}. {r['titulo']}", 6, total_chars, bold=True, viñeta=False)
        f = _parrafo(ws, f, r["texto"], 6, total_chars, viñeta=False)
    f += 1
    f = _seccion(ws, f, txt["limites"], 6)
    for c in g["limites"]:
        f = _parrafo(ws, f, c, 6, total_chars)
    _pagina(ws, horizontal=False)


def _por_mes(ws, A, txt, lang):
    T = A["anio"]
    ANCHO = 15.5
    n = len(txt["ficha"])
    for j in range(1, n + 1):
        ws.column_dimensions[get_column_letter(j)].width = ANCHO
    chars = ANCHO * n * 1.05
    f = _titulo(ws, A, txt, n)
    for m in A["meses"]:
        # Encabezado de la ficha
        ws.merge_cells(start_row=f, start_column=1, end_row=f, end_column=n - 2)
        c = ws.cell(row=f, column=1, value=f"{MESES[lang][m['i']]} {T}")
        c.font = _f(12, True, "FFFFFF"); c.fill = _fill(NAVY); c.alignment = Alignment(vertical="center", indent=1)
        ws.merge_cells(start_row=f, start_column=n - 1, end_row=f, end_column=n)
        c = ws.cell(row=f, column=n - 1, value=f"{txt['conf']}: {m['confianza_txt'].upper()}")
        c.font = _f(10, True, "FFFFFF"); c.fill = _fill(_color_conf(m["confianza"]))
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[f].height = 22
        f += 1
        if not m["sin_operacion"]:
            for j, (_, et, _) in enumerate(txt["ficha"], start=1):
                c = _celda(ws, f, j, et, h="center", fill=CLARO, color=C["text_mid"])
                c.font = _f(9, True, C["text_mid"])
            f += 1
            for j, (k, _, fmt) in enumerate(txt["ficha"], start=1):
                _celda(ws, f, j, m.get(k), fmt, bold=k == "proj", h="center")
            ws.row_dimensions[f].height = 18
            f += 1
        f = _parrafo(ws, f, m["comentario"], n, chars, viñeta=False)
        if m["razones_txt"]:
            ws.merge_cells(start_row=f, start_column=1, end_row=f, end_column=n)
            t = txt["razones"] + "; ".join(m["razones_txt"]) + "."
            c = ws.cell(row=f, column=1, value=t)
            c.font = _f(9, False, C["text_mid"], True); c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
            ws.row_dimensions[f].height = _alto(t, chars * 1.1)
            f += 1
        f += 1
    _pagina(ws, horizontal=False)


def _posicion(ws, A, txt, lang):
    P, T = A["posicion"], A["anio"]
    f = _titulo(ws, A, txt, len(txt["pos"]))
    _encabezado(ws, f, txt["pos"], [16, 11, 11, 9, 12, 14, 10, 9, 10, 11, 12, 11, 12, 10, 14, 20])
    ws.freeze_panes = ws.cell(row=f + 1, column=2)
    f += 1
    meses = {m["i"]: m for m in A["meses"]}
    for r in P["rows"]:
        m = meses[r["i"]]
        bg = GRIS if r["i"] % 2 else None
        vals = [(f"{MESES[lang][r['i']]} {T}", None), (r["cap"], F_RN), (r["otb_rn"], F_RN), (r["occ"], F_PCT),
                (r["adr"], F_USD), (r["otb_rev"], F_USD), (r["st_rn"], F_RN), (r["pace"], F_PCT), (r["bloqueos"], F_RN),
                (r["ly_rn"], F_RN), (txt["base_lab"][m["ly_base"]], None),
                (None if r["cerrado"] else max(0.0, r["proj"] - r["otb_rn"]), F_RN), (r["proj"], F_RN),
                (r["proj_occ"], F_PCT), (r["proj_rev"], F_USD), (txt["estados"].get(r.get("estado"), ""), None)]
        for j, (v, fmt) in enumerate(vals, start=1):
            _celda(ws, f, j, v, fmt, bold=j in (13, 15), fill=bg, h="left" if j in (1, 11, 16) else "right")
        f += 1
    t = P["tot"]
    tot = [txt["total"], t["cap"], t["otb_rn"], t["occ"], t["adr"], t["otb_rev"], t["st_rn"], t["pace"], t["bloqueos"],
           t["ly_rn"], "", max(0.0, t["proj"] - t["otb_rn"]), t["proj"], t["proj_occ"], t["proj_rev"], ""]
    fmts = [None, F_RN, F_RN, F_PCT, F_USD, F_USD, F_RN, F_PCT, F_RN, F_RN, None, F_RN, F_RN, F_PCT, F_USD, None]
    for j, (v, fmt) in enumerate(zip(tot, fmts), start=1):
        _celda(ws, f, j, v, fmt, bold=True, fill=CLARO, h="left" if j == 1 else "right")
    _pagina(ws, filas_titulo="4:4")


def _detalle(ws, A, txt, lang):
    T = A["anio"]
    f = _titulo(ws, A, txt, len(txt["det"]))
    _encabezado(ws, f, txt["det"], [16] + [11.5] * (len(txt["det"]) - 2) + [12])
    ws.freeze_panes = ws.cell(row=f + 1, column=2)
    f += 1
    claves = [("otb", F_RN), ("stly", F_RN), ("pace", F_PCT), ("proj", F_RN), ("proj_occ", F_PCT), ("proj_rev", F_USD),
              ("proj_vs_ly", None), ("rev_vs_ly", None), ("asegurado", F_PCT), ("falta", F_RN), ("ly_pick", F_RN),
              ("esfuerzo", F_PCT), ("por_semana", F_RN), ("ritmo_semana", "#,##0.0"), ("otb_adr", F_USD), ("st_adr", F_USD),
              ("tarifa_vs_stly", F_PCT), ("libres", F_RN), ("valor_5pp", F_USD), ("valor_tarifa_5", F_USD)]
    for m in A["meses"]:
        bg = GRIS if m["i"] % 2 else None
        _celda(ws, f, 1, f"{MESES[lang][m['i']]} {T}", h="left", fill=bg)
        for j, (k, fmt) in enumerate(claves, start=2):
            v = m.get(k)
            if k in ("proj_vs_ly", "rev_vs_ly"):
                v, fmt = (v - 1 if v is not None else None), "+0%;-0%;0%"
            _celda(ws, f, j, v, fmt, fill=bg, bold=k == "proj")
        c = _celda(ws, f, len(claves) + 2, m["confianza_txt"], h="center", fill=bg, bold=True, color=_color_conf(m["confianza"]))
        f += 1
    _pagina(ws, filas_titulo="4:4")


def _base(ws, A, txt, lang):
    B, T = A["base"]["anio"], A["anio"]
    st = (A["stly"] or {}).get("as_of") or "—"
    cab = [h.replace("{st}", st) for h in txt["base"]]
    f = _titulo(ws, A, txt, len(cab))
    _encabezado(ws, f, cab, [16, 11, 11, 10, 9, 14, 11, 12, 14, 11, 13, 12, 13, 12])
    ws.freeze_panes = ws.cell(row=f + 1, column=2)
    f += 1
    for x in A["base"]["filas"]:
        bg = GRIS if x["i"] % 2 else None
        est = txt["real"] if x["cerrado"] else txt["curso"] if x["abierto"] else txt["cerrado_h"]
        vs = x["rn"] / x["meta_rn"] - 1 if x["meta_rn"] else None
        vals = [(f"{MESES[lang][x['i']]} {B}", None), (est, None), (x["cap"], F_RN), (x["rn"], F_RN), (x["occ"], F_PCT),
                (x["rev"], F_USD), (x["adr"], F_USD), (x["meta_rn"], F_RN), (x["meta_rev"], F_USD), (vs, "+0%;-0%;0%"),
                (x["stly"], F_RN), (x["pick"], F_RN), (x["tasa"], F_PCT), (x["reconcilia"], F_PCT)]
        for j, (v, fmt) in enumerate(vals, start=1):
            _celda(ws, f, j, v, fmt, fill=bg, h="left" if j <= 2 else "right",
                   color=C["text_mid"] if (j == 2 and not x["cerrado"]) else "1A1A2E", bold=j == 13)
        f += 1
    t = A["base"]["tot"]
    vs = t["rn"] / t["meta_rn"] - 1 if t.get("meta_rn") else None
    tot = [(txt["total"], None), ("", None), (t["cap"], F_RN), (t["rn"], F_RN), (t["occ"], F_PCT), (t["rev"], F_USD),
           (t["adr"], F_USD), (t["meta_rn"], F_RN), (t["meta_rev"], F_USD), (vs, "+0%;-0%;0%"),
           (sum(x["stly"] or 0 for x in A["base"]["filas"]), F_RN), (sum(x["pick"] or 0 for x in A["base"]["filas"]), F_RN),
           (None, None), (t["reconcilia"], F_PCT)]
    for j, (v, fmt) in enumerate(tot, start=1):
        _celda(ws, f, j, v, fmt, bold=True, fill=CLARO, h="left" if j <= 2 else "right")
    f += 2
    f = _parrafo(ws, f, txt["nota_base"].replace("{T}", str(T)), len(cab), 170, viñeta=False)
    f += 1
    vent = A["base"]["ventana"]
    if vent:
        f = _seccion(ws, f, txt["ventana"].replace("{B}", str(B)), len(cab))
        from app.engine.pacing_tecnico import LEADS
        _encabezado(ws, f, [txt["base"][0]] + [txt["dias"].format(d=d) for d in LEADS])
        f += 1
        for v in vent:
            _celda(ws, f, 1, f"{MESES[lang][v['i']]} {B}", h="left")
            for j, d in enumerate(LEADS, start=2):
                p = v["puntos"].get(str(d))
                _celda(ws, f, j, p if p is not None else txt["nd"], F_PCT, h="right" if p is not None else "center",
                       color="1A1A2E" if p is not None else C["text_mid"])
            f += 1
    _pagina(ws, filas_titulo="4:4")


def _metodos(ws, A, txt, lang):
    T, esc = A["anio"], A["escenario"]
    f = _titulo(ws, A, txt, 6)
    _encabezado(ws, f, txt["metodos"], [26, 16, 16, 18, 14, 16])
    f += 1
    tot = {k: 0.0 for k in ("avail", "add", "mult")}; rev = 0.0
    for m in A["meses"]:
        bg = GRIS if m["i"] % 2 else None
        _celda(ws, f, 1, f"{MESES[lang][m['i']]} {T}", h="left", fill=bg)
        for j, k in enumerate(("avail", "add", "mult"), start=2):
            _celda(ws, f, j, m["metodos"][k], F_RN, fill=CLARO if k == esc else bg, bold=k == esc)
            tot[k] += m["metodos"][k]
        _celda(ws, f, 5, txt["met"][esc], h="center", fill=bg)
        _celda(ws, f, 6, m["metodos_rev"][esc], F_USD, fill=bg)
        rev += m["metodos_rev"][esc]
        f += 1
    _celda(ws, f, 1, txt["total"], h="left", bold=True, fill=CLARO)
    for j, k in enumerate(("avail", "add", "mult"), start=2):
        _celda(ws, f, j, tot[k], F_RN, bold=True, fill=CLARO)
    _celda(ws, f, 5, "", fill=CLARO); _celda(ws, f, 6, rev, F_USD, bold=True, fill=CLARO)
    f += 2
    f = _seccion(ws, f, txt["glosario"], 6)
    for k, v in txt["gl"]:
        ws.cell(row=f, column=1, value=k).font = _f(10, True, NAVY)
        ws.cell(row=f, column=1).alignment = Alignment(vertical="top", wrap_text=True)
        ws.merge_cells(start_row=f, start_column=2, end_row=f, end_column=6)
        c = ws.cell(row=f, column=2, value=v)
        c.font = _f(10); c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[f].height = _alto(v, 82)
        f += 1
    _pagina(ws, horizontal=False)


def construir_excel(A: dict, hotel: str) -> bytes:
    lang = A.get("lang", "es")
    txt0 = L[lang]
    T, B = A["anio"], A["base"]["anio"]
    txt = dict(txt0)
    txt["titulo"] = txt0["titulo"].format(hotel=hotel, T=T)
    txt["sub"] = txt0["sub"].format(corte=A["corte"], kind=txt0["kinds"].get(A["kind"], A["kind"]),
                                    met=txt0["met"][A["escenario"]], hoy=datetime.now().strftime("%Y-%m-%d"))
    wb = Workbook()
    hojas = [h.format(T=T, B=B) for h in txt0["hojas"]]
    fns = [_resumen, _por_mes, _posicion, _detalle, _base, _metodos]
    for k, (nombre, fn) in enumerate(zip(hojas, fns)):
        ws = wb.active if k == 0 else wb.create_sheet()
        ws.title = nombre[:31]
        fn(ws, A, txt, lang)
    wb.properties.title = txt["titulo"]
    wb.properties.creator = "FinPlan"
    return workbook_to_bytes(wb)
