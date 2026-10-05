# -*- coding: utf-8 -*-
"""El reporte de discrepancias del mayor, en Excel.

Tres hojas, en el orden en que el owner las usa:

1. **Discrepancias** — un renglon por caso, de mayor a menor severidad y monto.
   Trae una columna `Revisado` vacia: el reporte es la hoja de trabajo, no un
   informe para leer y archivar.
2. **Detalle** — cada linea de asiento de cada caso, con numero de asiento y
   fecha, para ir a buscarla en Integrity.
3. **Reglas** — que mira cada regla y por que. Sin esto el reporte pide
   confianza ciega; con esto el owner puede discutir una regla que no le sirva.
"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_HDR = PatternFill("solid", fgColor="1E2130")
_HDR_F = Font(color="FFFFFF", bold=True, size=10)
_SEV = {"alta": PatternFill("solid", fgColor="FFD6D6"),
        "media": PatternFill("solid", fgColor="FFF0CC"),
        "baja": PatternFill("solid", fgColor="EDEDED")}
_BORDE = Border(bottom=Side(style="thin", color="D0D0D0"))
_MONEDA = '#,##0.00;[Red]-#,##0.00'

QUE_MIRA = {
    "ARTICULO_CRUZA_CLASE": (
        "El mismo articulo quedo en costo de ventas (5xxx) y en gasto de operacion "
        "(7xxx). No puede ser los dos. Se senala el lado con menos lineas; si estan "
        "parejos se senalan los dos, porque ahi hay que decidir."),
    "ARTICULO_EN_VARIAS_CUENTAS": (
        "El mismo articulo en varias cuentas de la MISMA clase, con un reparto "
        "desparejo. Un reparto parejo no se avisa: huevos al restaurante y al "
        "comedor de empleados es a proposito."),
    "COSTO_EN_DEPTO_SIN_COSTO": (
        "Una cuenta de costo de ventas (5xxx) en un departamento que no vende un "
        "producto directo — Rooms, Admin, Ventas, Mantenimiento."),
    "PLANILLA_DEPTO_NO_COINCIDE": (
        "El departamento que dice el texto del asiento no es el de la cuenta. Es el "
        "error que aparecio en agosto: planilla del comedor de empleados rotulada "
        "FRONT DESK."),
    "PLANILLA_PUESTO_NO_COINCIDE": (
        "Un mismo codigo de puesto aparece con dos nombres distintos en el mes."),
    "CUENTA_GENERICA": (
        "Gasto cargado a una cuenta de cajon (7380 Miscellaneous). No es un error "
        "por si mismo: es donde se esconde uno."),
    "ALLOCATION_NO_NETEA": (
        "La cuenta 4999 reparte gasto entre departamentos y tiene que sumar cero a "
        "nivel de hotel. Si no suma cero, el reparto quedo a medias."),
}

_COLS = [
    ("Revisado", 10), ("Severidad", 10), ("Regla", 28), ("Cuenta", 28),
    ("Depto", 7), ("Concepto", 34), ("Lineas", 7), ("Monto CRC", 15),
    ("Monto USD", 13), ("Deberia ir en", 14), ("Por que", 96),
]
_COLS_DET = [
    ("Severidad", 10), ("Regla", 28), ("Concepto", 30), ("Cuenta", 28),
    ("Asiento", 9), ("Linea", 7), ("Fecha", 11), ("Origen", 8),
    ("Descripcion", 40), ("Referencia", 36), ("Monto CRC", 15), ("Monto USD", 13),
]


def _encabezado(ws, cols) -> None:
    for i, (nombre, ancho) in enumerate(cols, 1):
        c = ws.cell(row=1, column=i, value=nombre)
        c.fill, c.font = _HDR, _HDR_F
        c.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = "A2"


def construir(resumen, archivo: str = "") -> bytes:
    wb = Workbook()

    ws = wb.active
    ws.title = "Discrepancias"
    _encabezado(ws, _COLS)
    for h in resumen.hallazgos:
        ws.append(["", h.severidad.upper(), h.regla, h.cuenta, h.dept, h.concepto,
                   h.n_lineas, h.monto_crc, h.monto_usd, h.sugerencia, h.porque])
        f = ws.max_row
        for col in range(1, len(_COLS) + 1):
            ws.cell(row=f, column=col).border = _BORDE
        ws.cell(row=f, column=2).fill = _SEV[h.severidad]
        ws.cell(row=f, column=8).number_format = _MONEDA
        ws.cell(row=f, column=9).number_format = _MONEDA
        ws.cell(row=f, column=11).alignment = Alignment(wrap_text=True, vertical="top")
    ws.auto_filter.ref = "A1:K{}".format(max(ws.max_row, 1))

    det = wb.create_sheet("Detalle")
    _encabezado(det, _COLS_DET)
    for h in resumen.hallazgos:
        for l in h.lineas:
            det.append([h.severidad.upper(), h.regla, h.concepto, l["cuenta"],
                        l["asiento"], l["linea"], l["fecha"], l["origen"],
                        l["descripcion"], l["referencia"], l["monto_crc"],
                        l["monto_usd"]])
            f = det.max_row
            det.cell(row=f, column=1).fill = _SEV[h.severidad]
            det.cell(row=f, column=11).number_format = _MONEDA
            det.cell(row=f, column=12).number_format = _MONEDA
    det.auto_filter.ref = "A1:L{}".format(max(det.max_row, 1))

    rg = wb.create_sheet("Reglas")
    rg.append(["Auditoria del detalle del mayor"])
    rg["A1"].font = Font(bold=True, size=13)
    rg.append([])
    for etiqueta, valor in (
            ("Periodo", resumen.periodo),
            ("Archivo", archivo),
            ("Lineas del archivo", resumen.lineas_del_archivo),
            ("Lineas revisadas (clase 4 en adelante)", resumen.lineas_revisadas),
            ("Casos a revisar", len(resumen.hallazgos)),
            ("Lineas senaladas", resumen.lineas_senaladas),
            ("Monto senalado (CRC)", resumen.monto_en_revision_crc)):
        rg.append([etiqueta, valor])
    rg.append([])
    rg.append(["Regla", "Casos", "Que mira"])
    for c in ("A", "B", "C"):
        rg["{}{}".format(c, rg.max_row)].fill = _HDR
        rg["{}{}".format(c, rg.max_row)].font = _HDR_F
    for regla, n in sorted(resumen.por_regla.items(), key=lambda x: -x[1]):
        rg.append([regla, n, QUE_MIRA.get(regla, "")])
        rg.cell(row=rg.max_row, column=3).alignment = Alignment(wrap_text=True,
                                                               vertical="top")
    rg.append([])
    rg.append(["Esto no decide si un asiento esta bien: senala los que no se parecen "
               "al resto, para que la revision empiece por ahi."])
    for col, ancho in (("A", 38), ("B", 9), ("C", 110)):
        rg.column_dimensions[col].width = ancho

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
