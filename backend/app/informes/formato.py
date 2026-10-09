# -*- coding: utf-8 -*-
"""El formato del informe: estilos, tablas y los bloques que se repiten."""
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

AZUL   = RGBColor(0x1F, 0x3B, 0x63)
GRIS   = RGBColor(0x55, 0x5A, 0x66)
ROJO   = RGBColor(0xB3, 0x26, 0x1E)
VERDE  = RGBColor(0x1B, 0x6B, 0x4F)
NEGRO  = RGBColor(0x1A, 0x1A, 0x1A)

def doc_nuevo():
    d = Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(1.9)
        s.top_margin = s.bottom_margin = Cm(1.9)
    n = d.styles["Normal"]
    n.font.name = "Calibri"; n.font.size = Pt(10)
    n.paragraph_format.space_after = Pt(5)
    n.paragraph_format.line_spacing = 1.08
    return d

def _sombra(celda, hexa):
    el = OxmlElement("w:shd"); el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hexa); celda._tc.get_or_add_tcPr().append(el)

def h1(d, txt, num=None):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(16); p.paragraph_format.space_after = Pt(4)
    r = p.add_run((f"{num}. " if num else "") + txt.upper())
    r.bold = True; r.font.size = Pt(13); r.font.color.rgb = AZUL
    b = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom")
    bt.set(qn("w:val"), "single"); bt.set(qn("w:sz"), "10")
    bt.set(qn("w:color"), "1F3B63"); b.append(bt)
    p._p.get_or_add_pPr().append(b)
    return p

def h2(d, txt):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(11); p.paragraph_format.space_after = Pt(3)
    r = p.add_run(txt); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = AZUL
    return p

def h3(d, txt):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(txt); r.bold = True; r.font.size = Pt(10); r.font.color.rgb = NEGRO
    return p

def par(d, txt, size=10, italic=False, color=None, space=5):
    p = d.add_paragraph()
    p.paragraph_format.space_after = Pt(space)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(txt); r.font.size = Pt(size); r.italic = italic
    if color is not None: r.font.color.rgb = color
    return p

def bullet(d, txt, size=10, nivel=0):
    p = d.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.6 + 0.5 * nivel)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(txt); r.font.size = Pt(size)
    return p

def nota(d, titulo, txt):
    """El recuadro gris de advertencia o de contexto."""
    t = d.add_table(rows=1, cols=1); t.style = "Table Grid"
    c = t.cell(0, 0); _sombra(c, "F2F4F7")
    c.paragraphs[0].paragraph_format.space_after = Pt(2)
    r = c.paragraphs[0].add_run(titulo); r.bold = True; r.font.size = Pt(9)
    r.font.color.rgb = AZUL
    p = c.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rr = p.add_run(txt); rr.font.size = Pt(9); rr.font.color.rgb = GRIS
    d.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def tabla(d, cabeceras, filas, anchos=None, nota_pie=None, size=8.5):
    """`filas`: lista de listas. Un elemento puede ser (texto, estilo) donde
    estilo es 'tot' (negrita + fondo), 'sub' (negrita), 'neg'/'pos' (color)."""
    t = d.add_table(rows=1, cols=len(cabeceras))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for i, h in enumerate(cabeceras):
        c = t.rows[0].cells[i]; _sombra(c, "1F3B63")
        c.paragraphs[0].paragraph_format.space_after = Pt(0)
        c.paragraphs[0].alignment = (WD_ALIGN_PARAGRAPH.LEFT if i == 0
                                     else WD_ALIGN_PARAGRAPH.RIGHT)
        r = c.paragraphs[0].add_run(str(h)); r.bold = True; r.font.size = Pt(size)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for n, fila in enumerate(filas):
        celdas = t.add_row().cells
        for i, v in enumerate(fila):
            est = None
            # Desde JSON una celda con estilo llega como LISTA, no como tupla.
            if isinstance(v, (tuple, list)) and len(v) == 2: v, est = v
            cel = celdas[i]
            cel.paragraphs[0].paragraph_format.space_after = Pt(0)
            cel.paragraphs[0].alignment = (WD_ALIGN_PARAGRAPH.LEFT if i == 0
                                           else WD_ALIGN_PARAGRAPH.RIGHT)
            r = cel.paragraphs[0].add_run("" if v is None else str(v))
            r.font.size = Pt(size)
            if est in ("tot", "sub"): r.bold = True
            if est == "tot": _sombra(cel, "E8ECF3")
            if est == "sec": r.bold = True; _sombra(cel, "D7DEEA")
            if est == "neg": r.font.color.rgb = ROJO
            if est == "pos": r.font.color.rgb = VERDE
            if est == "gris": r.font.color.rgb = GRIS
        if n % 2 == 1 and not any(isinstance(x, tuple) and x[1] in ("tot","sec")
                                  for x in fila):
            for cel in celdas: _sombra(cel, "FAFBFD")
    if nota_pie:
        p = d.add_paragraph(); p.paragraph_format.space_after = Pt(8)
        r = p.add_run(nota_pie); r.font.size = Pt(8); r.italic = True
        r.font.color.rgb = GRIS
    else:
        d.add_paragraph().paragraph_format.space_after = Pt(4)
    return t

# ── formato de numeros ─────────────────────────────────────────────────────
def us(x, dec=2):
    if x is None: return "—"
    return f"{x:,.{dec}f}" if abs(x) > 0.0049 else "—"

def pc(x, dec=1):
    if x is None: return "—"
    return f"{x*100:,.{dec}f}%"

def var(a, b):
    """La variacion y su color segun si favorece o no. `b` es la referencia."""
    if a is None or b is None: return ("—", None)
    v = a - b
    return (us(v), "neg" if v < -0.005 else ("pos" if v > 0.005 else "gris"))

def var_gasto(a, b):
    """En gasto, MAS es desfavorable: el color se invierte."""
    if a is None or b is None: return ("—", None)
    v = a - b
    return (us(v), "neg" if v > 0.005 else ("pos" if v < -0.005 else "gris"))

def pct_var(a, b):
    if a is None or b is None or abs(b) < 0.005: return "—"
    return f"{(a-b)/abs(b)*100:,.1f}%"
